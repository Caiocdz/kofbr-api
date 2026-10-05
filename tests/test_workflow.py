import io
from copy import deepcopy
import openpyxl
import pytest
from app import create_app
from workflow import importer, service, store

HEADERS = ['Unidade', 'Data', 'Linha', 'Equipamento', 'Tipo de falha', 'Observações', 'Minutos de parada']


def spreadsheet(rows, headers=HEADERS):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Apontamentos'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    data = io.BytesIO()
    workbook.save(data)
    data.seek(0)
    return data


@pytest.fixture
def client(tmp_path, monkeypatch):
    from workflow import learning
    monkeypatch.setenv('KOFBR_DATA_DIR', str(tmp_path))
    learning.reset()
    # Nos testes o treino é chamado explicitamente (sem threads em segundo plano).
    monkeypatch.setattr(learning, 'retrain_in_background', lambda build: None)
    monkeypatch.setenv('KOFBR_WORKFLOW_DB', 'sqlite')
    # Sem pasta de treino o quadro agrupa só por semelhança (o teste do ML liga a pasta própria).
    from workflow import apontamentos_ml
    monkeypatch.setenv('KOFBR_ML_TRAIN_DIR', str(tmp_path / 'sem-treino'))
    monkeypatch.setattr(apontamentos_ml, 'retrain_soon', lambda delay=20.0: None)
    apontamentos_ml._state.update(model=None, info=None, fingerprint=None, training=False)
    app = create_app()
    app.config['TESTING'] = True
    with app.test_client() as client:
        yield client


def upload(client, rows, filename='apontamentos.xlsx'):
    return client.post('/api/workspace/import', data={'file': (spreadsheet(rows), filename)})


def action(client, board, verb, **kwargs):
    return client.post(f'/api/workspace/analyses/{board["id"]}/actions', json={'action': verb, 'revision': board['revision'], **kwargs})


def validate_all(client, board):
    for column in board['columns']:
        result = action(client, board, 'validate_column', column_id=column['id'])
        assert result.status_code == 200, result.json
        board = result.json
    return board


def test_group_unlimited_and_keep_context():
    rows = [['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'Rolamento quebrado', 5]] * 1307
    rows += [['U2', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'Rolamento quebrado', 5]]
    rows += [['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'Sem rolamento quebrado', 5]]
    records, _ = importer.parse_file(spreadsheet(rows).getvalue(), 'base.xlsx')
    columns, cards = importer.group_records(records)
    assert len(columns) == 2
    assert len(cards) == 3
    assert sorted(len(c['record_ids']) for c in cards) == [1, 1, 1307]
    assert records[0]['data_inicio'] == '2026-09-30'


def test_full_board_persistence_validation_hold_move_and_source(client):
    rows = [['U1', '30/09/2026', 'L1', 'Bomba', 'Mecânica', 'Vazamento no selo', 10],
            ['U1', '30/09/2026', 'L1', 'Bomba', 'Mecânica', 'Vazamento no selo', 20],
            ['U1', '29/09/2026', 'L1', 'Motor', 'Elétrica', 'Motor queimado', 40]]
    response = upload(client, rows)
    assert response.status_code == 201
    board = response.json
    analysis_id = board['id']
    assert board['records_count'] == 3 and board['groups_count'] == 2
    assert not board['ready']
    assert client.post(f'/api/workspace/analyses/{analysis_id}/finish').status_code == 409
    assert client.get(f'/api/workspace/analytics?ids={analysis_id}').status_code == 400
    motor = next(c for c in board['columns'] if c['name'] == 'Motor')
    board = action(client, board, 'validate_column', column_id=motor['id']).json
    pump = next(c for c in board['cards'] if c['count'] == 2)
    assert not pump['validated']
    # Ao entrar em uma coluna validada, o card herda a validação.
    board = action(client, board, 'move', card_id=pump['id'], column_id=motor['id']).json
    assert board['ready']
    stale = action(client, {**board, 'revision': 1}, 'hold', card_id=pump['id'])
    assert stale.status_code == 409
    # O sino torna o card pendente e bloqueia o dashboard inclusive pela API.
    board = action(client, board, 'hold', card_id=pump['id']).json
    assert board['held_count'] == 1 and not board['ready']
    assert client.post(f'/api/workspace/analyses/{analysis_id}/finish').status_code == 409
    board = action(client, board, 'move', card_id=pump['id'], column_id=motor['id']).json
    assert board['held_count'] == 0 and board['ready']
    # Não é possível excluir coluna ocupada.
    assert action(client, board, 'delete_column', column_id=motor['id']).status_code == 400
    empty = next(c for c in board['columns'] if c['name'] == 'Bomba')
    board = action(client, board, 'delete_column', column_id=empty['id']).json
    assert len(board['columns']) == 1
    reloaded = client.get(f'/api/workspace/analyses/{analysis_id}').json
    assert reloaded['ready'] and reloaded['revision'] == board['revision']
    assert client.post(f'/api/workspace/analyses/{analysis_id}/finish').status_code == 200
    data = client.get(f'/api/workspace/analytics?ids={analysis_id}&from=2026-09-30&to=2026-09-30').json
    assert data['metrics']['count'] == 2
    assert data['metrics']['minutes'] == 30
    assert data['machines'][0]['name'] == 'Motor'  # mudança refletida nos gráficos
    detail = client.get(f'/api/workspace/analyses/{analysis_id}/cards/{pump["id"]}').json
    assert all(r['equipamento'] == 'Bomba' for r in detail['records'])  # fonte original preservada
    source = client.get(f'/api/workspace/analyses/{analysis_id}/source')
    assert source.status_code == 200 and source.data[:2] == b'PK'
    assert client.get('/api/workspace/analyses').json[0]['days'] == ['2026-09-29', '2026-09-30']


def test_duplicate_content_and_invalid_sheet(client):
    rows = [['U1', '2026-09-30', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 4]]
    first = upload(client, rows).json
    duplicate = upload(client, rows, 'renomeada.xlsx')
    assert duplicate.status_code == 409 and duplicate.json['existing_id'] == first['id']
    assert len(client.get('/api/workspace/analyses').json) == 1
    for field, value in [(1, '30/02/2026'), (6, -1), (6, 'abc')]:
        wrong = deepcopy(rows)
        wrong[0][field] = value
        result = upload(client, wrong)
        assert result.status_code == 400, result.json
        assert 'Linha 2' in result.json['erro']
    result = client.post('/api/workspace/import', data={'file': (spreadsheet([['x']], ['Nome']), 'errada.xlsx')})
    assert result.status_code == 400 and 'Cabeçalhos' in result.json['erro']
    assert len(client.get('/api/workspace/analyses').json) == 1


def test_blank_cells_import_anyway(client):
    rows = [['U1', '2026-09-30', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 4],
            ['U1', '2026-09-30', 'L1', 'Bomba', 'Mecânica', 'Vazamento', None],
            ['U1', '2026-09-30', None, None, 'Mecânica', 'Vazamento', 3],
            ['U1', '2026-09-30', 'L1', None, 'Mecânica', 'Vazamento', 2],
            ['U1', None, 'L1', 'Bomba', 'Mecânica', 'Vazamento', 1],
            [None, '2026-09-30', 'L1', 'Bomba', None, None, 5]]
    result = upload(client, rows)
    assert result.status_code == 201, result.json
    board = result.json
    doc = store.get(board['id'])
    assert len(doc['records']) == 5
    assert [r['minutos_parada'] for r in doc['records']][:2] == [4, 0]
    loose = [r for r in doc['records'] if r.get('sem_contexto')]
    assert len(loose) == 2 and {r['equipamento'] for r in loose} == {'Não informado'}
    # Mesmo texto, mas sem contexto conhecido: cada um fica em card próprio.
    for r in loose:
        assert any(c['record_ids'] == [r['id']] for c in doc['cards'])
    warnings = ' '.join(doc['warnings'])
    assert 'sem data' in warnings and 'sem minutos' in warnings and 'sem linha ou equipamento' in warnings


def test_compare_missing_machines_zero_baseline_and_filters(client):
    rows = [
        ['U1', '2026-09-01', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 10],
        ['U1', '2026-09-01', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 30],
        ['U1', '2026-09-01', 'L1', 'Motor', 'Elétrica', 'Falha', 50],
        ['U1', '2026-09-02', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 20],
        ['U1', '2026-09-02', 'L1', 'Esteira', 'Mecânica', 'Trava', 80],
        ['U2', '2026-09-02', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 10],
    ]
    board = validate_all(client, upload(client, rows).json)
    result = client.get('/api/workspace/compare?from=2026-09-02&to=2026-09-02&unit=U1').json
    changes = {m['name']: m for m in result['changes']}
    assert changes['Motor']['percentage'] == -100
    assert changes['Bomba']['percentage'] == -50
    assert changes['Esteira']['percentage'] is None
    assert result['a']['metrics']['count'] == 3
    assert result['b']['metrics']['count'] == 2
    assert result['a']['metrics']['mttr'] == 30
    assert result['b']['metrics']['mttr'] == 50
    metrics = client.get('/api/workspace/analytics?unit=U2&line=L1&failure=Mecânica').json['metrics']
    assert metrics['count'] == 1 and metrics['minutes'] == 10
    assert client.get('/api/workspace/analytics?from=2026-10-01&to=2026-09-01').status_code == 400
    assert client.get('/api/workspace/compare?from=bad&to=bad').status_code == 400
    assert set(client.get('/api/workspace/filters').json['units']) == {'U1', 'U2'}
    for fmt, signature in [('pdf', b'%PDF'), ('png', b'\x89PNG')]:
        for chart in ['pareto', 'jackknife']:
            result = client.get(f'/api/workspace/export?chart={chart}&format={fmt}&ids={board["id"]}')
            assert result.status_code == 200
            assert result.data.startswith(signature)


def test_add_split_and_reopen_validation(client):
    board = upload(client, [['U', '2026-01-01', 'L', 'M', 'Tipo', 'Falha', 10]]*3).json
    board = validate_all(client, board)
    card = board['cards'][0]
    board = action(client, board, 'add_column', name='Nova peça', unit='U', line='L').json
    new = board['columns'][-1]
    board = action(client, board, 'move', card_id=card['id'], column_id=new['id']).json
    assert not board['ready']  # destino não validado pede revisão
    board = action(client, board, 'split_card', card_id=card['id']).json
    assert len(board['cards']) == 3
    assert sum(c['count'] for c in board['cards']) == 3
    board = validate_all(client, board)
    assert board['ready']
    board = action(client, board, 'rename_card', card_id=board['cards'][0]['id'], name='Nova descrição').json
    assert not board['ready']
    # Uma segunda inicialização lê o mesmo histórico.
    store.initialize()
    assert store.get(board['id'])['revision'] == board['revision']


def test_validate_all_requires_human_confirmation(client):
    rows = [['U1', '30/09/2026', 'L1', 'Bomba', 'Mecânica', 'Vazamento no selo', 10],
            ['U1', '30/09/2026', 'L1', 'Motor', 'Elétrica', 'Motor queimado', 40]]
    board = upload(client, rows).json
    refused = action(client, board, 'validate_all')
    assert refused.status_code == 400 and 'CONFIRMAR' in refused.json['erro']
    held = action(client, board, 'hold', card_id=board['cards'][0]['id']).json
    blocked = action(client, held, 'validate_all', confirm='CONFIRMAR')
    assert blocked.status_code == 400
    moved = action(client, held, 'move', card_id=held['cards'][0]['id'], column_id=held['columns'][0]['id']).json
    done = action(client, moved, 'validate_all', confirm='confirmar').json
    assert done['ready'] and all(c['validated'] for c in done['columns'])


def test_category_filter_and_excel_summary(client):
    rows = [['U1', '30/09/2026', 'L1', 'Bomba', 'Mecânica', 'Vazamento', 50]] * 3
    rows += [['U1', '30/09/2026', 'L1', 'Motor', 'Elétrica', 'Queimado', 5]]
    rows += [['U1', '30/09/2026', 'L1', 'Esteira', 'Mecânica', 'Correia', 90]]
    board = validate_all(client, upload(client, rows).json)
    full = client.get(f'/api/workspace/analytics?ids={board["id"]}').json
    # Medianas: 1 ocorrência e 50 min; valores iguais ao corte ficam no lado alto.
    assert full['categories'] == {'Crítico-crônico': 2, 'Crítico': 0, 'Crônico': 1, 'Conforto': 0}
    assert len(full['rows']) == 3
    only = client.get(f'/api/workspace/analytics?ids={board["id"]}&category=Crítico-crônico').json
    assert {m['name'] for m in only['machines']} == {'Bomba', 'Esteira'}
    assert only['metrics']['count'] == 4 and only['q_threshold'] == full['q_threshold']
    comfort = client.get(f'/api/workspace/analytics?ids={board["id"]}&category=Crônico').json
    assert [m['name'] for m in comfort['machines']] == ['Motor'] and comfort['machines'][0]['category'] == 'Crônico'
    assert client.get(f'/api/workspace/analytics?ids={board["id"]}&category=Inventada').status_code == 400
    xlsx = client.get(f'/api/workspace/export/xlsx?ids={board["id"]}&from=2026-09-30&to=2026-09-30')
    assert xlsx.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(xlsx.data))
    assert book.sheetnames == ['Apontamentos Resumidos KOFBR', 'Tabela de máquinas', 'Tabela de falhas', 'Falhas por máquina']
    assert [c.value for c in book['Tabela de máquinas'][1]] == [None, 'EQUIPAMENTOS', 'T (min)', 'Q', 'MTTR', 'T.T.A', '%', '80%', 'Categoria']
    sheet = book['Apontamentos Resumidos KOFBR']
    assert [c.value for c in sheet[1]][:7] == HEADERS
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    # Três apontamentos iguais da Bomba viram uma única linha com os minutos somados.
    bomba = next(r for r in rows if r[3] == 'Bomba')
    assert len(rows) == 3 and bomba[6] == 150 and bomba[8] == 3 and bomba[5] == 'Vazamento'
    assert bomba[7] == 'SEM MODO DE FALHA IDENTIFICADO'  # 'Vazamento' não cita o item
    machines = book['Tabela de máquinas']
    assert machines['B2'].value == 'L1_Bomba' and machines.max_row >= 41 and machines['A41'].value == 40


def test_excel_keeps_sap_layout_and_merges_hourly_rows(client):
    headers = ['Centro', 'Data Inicio Real', 'Linha', 'Tipo de Parada', 'Material', 'Ordem', 'Descrição do Material',
               'Turno', 'Intervalo', 'chave do parada', 'Subchave de Parada', 'chave 1 de parada', 'Observações',
               'Pros. Efi. Perdid.', 'Ptos. Acumilados', 'Caixas Produzidas', 'Total minutos', 'Minutos de paradas']
    base = ['B0AC', '02/01/2026', 'LINHA003', 'P.EQ.LINHA', '92601', '11275225', 'CRYST PET 5L', 2]
    rows = [base + [f'{h:02d}:00-{h+1:02d}:00', 'PALETIZADOR 2 IEF', '', 'SISTEMA ESTRUTURA',
                    'FALHA NO EMPURRADOR DE CAMADAS 30008216826', 0.0035, 0.0035 * (h - 7), 0, 60, 60] for h in (8, 9, 10)]
    rows.append(base + ['13:00-14:00', 'ENCHEDORA BRH LI01', '', 'Sistema de enchimento', 'TROCA DA VALVULA', 0.0018, 0.02, 0, 31.2, 31.2])
    response = client.post('/api/workspace/import', data={'file': (spreadsheet(rows, headers), 'sap.xlsx')})
    board = validate_all(client, response.json)
    xlsx = client.get(f'/api/workspace/export/xlsx?ids={board["id"]}&from=2026-01-02&to=2026-01-02')
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.data)).worksheets[0]
    assert [c.value for c in sheet[1]][:18] == headers
    lines = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(lines) == 2
    pal = next(r for r in lines if r[9] == 'PALETIZADOR 2 IEF')
    assert pal[8] == '08:00-11:00' and pal[17] == 180 and pal[16] == 180 and pal[7] == 2 and pal[19] == 3 and pal[15] == 0 and pal[18] == 'FALHA DE EMPURRADOR'
    assert pal[12] == 'FALHA NO EMPURRADOR DE CAMADAS 30008216826'
    assert sheet['J1'].fill.fgColor.rgb.endswith('FFFF00') and sheet['M1'].fill.fgColor.rgb.endswith('C00000')


DOC_SENSOR = ['ATUADOR DO SENSOR TRAVADO ORDEM: 30008223804',
              'FALHA NO SENSOR DA HASTE DE PEGADA DE CAMADAS ( 30008224330 )',
              'FALHA NO SENSOR DE SAIDA DE PALETES ( 30008224349 )',
              'SENSOR DE PRESENÇA DE PALETES FOI MONTADO FORA DA POSIÇÃO NO PCM, OCASIONANDO FALHA NA BARREIRA DE SAIDA DA ENVOLVEDORA.',
              'SENSOR DA CENTRAGEM ESQUERDA DA MESA DE FORMAÇÃO DE CAMADA SAINDO FORA DE POSIÇÃO.']
DOC_VALVE = ['FALHA DE ENCHIMENTO, EXCESSO DE NIVEL BAIXO, VALVULA Nº39 QUE ESTÁ COM JAMPER COM A Nº36',
             'FALHA DE ENCHIMENTO, DIFICULTANDO DEIXAR ENCHEDORA NA VELOCIDADE NOMINAL, IDENTIFICADO VÁLVULA 41 TENDO MAIOR IMPACTO',
             'VALVULA Nº41 EM VELOCIDADE NOMINAL COM EXCESSO DE NIVEL BAIXO, VELOCIDADE REDUZIDA PARA MINIMIZAR PERDAS',
             'FALHA DE ENCHIMENTO NA VALVULA Nº36, EXCESSO DE NIVEL BAIXO',
             'VELOCIDADE REDUZIDA, FALHA DE ENCHIMENTO NA VALVULA Nº36, EXCESSO DE NIVEL BAIXO']


def test_failure_classification_matches_project_document():
    from workflow.classify import classify, default_catalog
    cat = default_catalog()
    assert {classify(t, 'PALETIZADOR 2 IEF', cat) for t in DOC_SENSOR} == {'FALHA DE SENSOR'}
    assert {classify(t, 'ENCHEDORA BRH LI01', cat) for t in DOC_VALVE} == {'FALHA DE VÁLVULA DE ENCHIMENTO'}
    assert classify('QUEBRA DE EIXO FUSIVEL E EIXO DE ARVORE', '', cat) == 'FALHA DE EIXO'
    assert classify('PARAFUSO ESPANADO NA ESTRELA', '', cat) == 'FALHA DE PARAFUSO'
    # O mesmo modo de falha escrito de jeitos diferentes cai na mesma classe (slide do funil).
    assert {classify(t, '', cat) for t in ['ROLAMENTO QUEBRADO', 'rolamento estourado', 'ROLAMNETO ESTOURADO 30008216805',
                                            'troca de rolamentos do motor']} == {'FALHA DE ROLAMENTO'}
    assert classify('ERRO DRIVE ALIMENTADOR ( QUEBRA DA CORREIA E ROLAMENTOS)', '', cat) == 'FALHA DE DRIVE'
    assert classify('ORDEM: 30008457970', '', cat) == 'SEM DESCRIÇÃO'
    assert classify('', '', cat) == 'SEM DESCRIÇÃO'
    assert classify('FALHA NA PALETIZADORA', '', cat) == 'SEM MODO DE FALHA IDENTIFICADO'


def test_failure_table_manual_class_machine_filter_and_catalog(client):
    rows = [['U1', '13/09/2026', 'LINHA003', 'PALETIZADOR 2 IEF', 'P.EQ.LINHA', t, 10] for t in DOC_SENSOR]
    rows += [['U1', '13/09/2026', 'LINHA001', 'ENCHEDORA BRH LI01', 'P.EQ.LINHA', t, 20] for t in DOC_VALVE]
    rows += [['U1', '13/09/2026', 'LINHA001', 'ENCHEDORA BRH LI01', 'P.EQ.LINHA', 'PARADA GERAL DA MAQUINA', 5]]
    board = upload(client, rows).json
    assert board['unclassified_count'] == 1
    vague = next(c for c in board['cards'] if c['class'] == 'SEM MODO DE FALHA IDENTIFICADO')
    board = action(client, board, 'set_class', card_id=vague['id'], failure_class='Falha de motor').json
    assert next(c for c in board['cards'] if c['id'] == vague['id'])['class'] == 'FALHA DE MOTOR'
    assert board['unclassified_count'] == 0
    board = validate_all(client, board)
    data = client.get(f'/api/workspace/analytics?ids={board["id"]}').json
    failures = {f['key']: f for f in data['failures']}
    assert failures['FALHA DE VÁLVULA DE ENCHIMENTO']['count'] == 5 and failures['FALHA DE SENSOR']['count'] == 5
    assert failures['FALHA DE MOTOR']['minutes'] == 5 and data['machines'][0]['short'] == 'L01_ENCHEDORA BRH LI01'
    only = client.get(f'/api/workspace/analytics?ids={board["id"]}&machine=PALETIZADOR 2 IEF').json
    assert only['metrics']['count'] == 5 and [f['key'] for f in only['failures']] == ['FALHA DE SENSOR']
    assert 'PALETIZADOR 2 IEF' in client.get('/api/workspace/filters').json['machines']
    # Catálogo editável: um novo termo passa a valer para os indicadores.
    catalog = client.get('/api/workspace/catalog').json['catalog']
    catalog['classes'].insert(0, {'label': 'PARADA GERAL', 'kind': 'regra', 'terms': ['PARADA GERAL']})
    assert client.put('/api/workspace/catalog', json=catalog).status_code == 200
    test = client.post('/api/workspace/catalog/test', json={'texts': ['PARADA GERAL DA MAQUINA']}).json
    assert test['results'][0]['class'] == 'PARADA GERAL'
    assert client.delete('/api/workspace/catalog').json['custom'] is False


SAP_HEADERS = ['Centro', 'Data Inicio Real', 'Linha', 'Tipo de Parada', 'Intervalo', 'chave do parada', 'Observações',
               'Minutos de paradas']


def test_one_failure_per_order_even_when_sap_slices_hours(client):
    # Uma quebra de 3 h (mesma O.S.) vira 3 linhas no SAP, mas é UMA falha.
    rows = [['B0AC', '02/01/2026', 'LINHA003', 'P.EQ.LINHA', f'{h:02d}:00-{h+1:02d}:00', 'PALETIZADOR',
             'ROLAMENTO QUEBRADO 30008216826', 60] for h in (8, 9, 10)]
    # Sem O.S.: intervalos encostados e mesmo relato também são a mesma falha; separados, não.
    rows += [['B0AC', '02/01/2026', 'LINHA003', 'P.EQ.LINHA', iv, 'PALETIZADOR', 'ROLAMENTO ESTOURADO', 30]
             for iv in ('13:00-14:00', '14:00-15:00', '18:00-19:00')]
    board = validate_all(client, client.post('/api/workspace/import', data={'file': (spreadsheet(rows, SAP_HEADERS), 'f.xlsx')}).json)
    assert board['automation']['lines'] == 6 and board['automation']['failures'] == 3
    events = client.get(f'/api/workspace/analytics?ids={board["id"]}').json
    lines = client.get(f'/api/workspace/analytics?ids={board["id"]}&count=lines').json
    assert events['count_mode'] == 'events' and events['metrics']['count'] == 3 and events['metrics']['lines'] == 6
    assert lines['metrics']['count'] == 6
    machine = events['machines'][0]
    assert machine['count'] == 3 and machine['lines'] == 6 and machine['mttr'] == 90
    assert [f['key'] for f in events['failures']] == ['FALHA DE ROLAMENTO'] and events['failures'][0]['count'] == 3
    assert client.get(f'/api/workspace/analytics?ids={board["id"]}&count=xx').status_code == 400


def test_confidence_memory_and_safe_batch(client):
    rows = [['U1', '13/09/2026', 'L1', 'ENCHEDORA', 'P.EQ.LINHA', 'SENSOR COM DEFEITO', 10],
            ['U1', '13/09/2026', 'L1', 'ENCHEDORA', 'P.EQ.LINHA', 'MAQUINA PAROU SOZINHA', 10]]
    board = upload(client, rows).json
    cards = {c['sample']: c for c in board['cards']}
    assert cards['SENSOR COM DEFEITO']['confidence'] == 'alta' and 'sensor' in cards['SENSOR COM DEFEITO']['reason']
    vague = cards['MAQUINA PAROU SOZINHA']
    assert vague['confidence'] == 'baixa'
    # Lote seguro exige confirmação e só valida o que tem confiança alta.
    assert action(client, board, 'validate_confident', card_ids=[cards['SENSOR COM DEFEITO']['id']]).status_code == 400
    board = action(client, board, 'validate_confident', confirm=True, card_ids=[cards['SENSOR COM DEFEITO']['id']]).json
    assert sum(c['validated'] for c in board['cards']) == 1
    board = action(client, board, 'set_class', card_id=vague['id'], failure_class='FALHA DE MOTOR ELÉTRICO').json
    assert next(c for c in board['cards'] if c['id'] == vague['id'])['confidence'] == 'manual'
    # Memória: o mesmo relato (mesmo com outra O.S.) chega classificado na próxima importação.
    assert client.get('/api/workspace/memory').json['total'] == 1
    again = upload(client, [['U1', '14/09/2026', 'L1', 'ENCHEDORA', 'P.EQ.LINHA', 'Máquina parou sozinha 30008216826', 7]], 'dia2.xlsx').json
    card = again['cards'][0]
    assert card['class'] == 'FALHA DE MOTOR ELÉTRICO' and card['confidence'] == 'alta' and card['source'] == 'memoria'


# ---------- Gerar planilha de apontamentos (ML) ----------
def ml_folder(tmp_path):
    folder = tmp_path / 'treino'
    folder.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['Observações Operação', 'Classificação Manual Analista'])
    parts = {'Sensor danificado': 'sensor', 'Rolamento quebrado': 'rolamento', 'Correia rompida': 'correia'}
    for label, word in parts.items():
        for i in range(25):
            sheet.append([f'{word} com defeito na posicao {i} da maquina', label])
    sheet.append(['sensor com defeito na posicao 0 da maquina', 'Sensor danificado'])   # duplicada
    sheet.append(['sensor com defeito na posicao 1 da maquina', 'Sensores danificados'])  # plural → unido
    sheet.append(['sensor com defeito na posicao 2 da maquina', 'Rolamento quebrado'])   # conflito (maioria vence)
    sheet.append([None, 'Sensor danificado'])                                             # faltante
    sheet.append(['correia gasta', None])                                                 # faltante
    sheet.append(['xxxxxx', 'Correia rompida'])                                           # sem conteúdo
    sheet.append(['motor parou', 'Sem detalhes'])                                         # rótulo sem informação
    workbook.save(folder / 'treino.xlsx')
    return folder


@pytest.fixture
def ml(client, tmp_path, monkeypatch):
    from workflow import apontamentos_ml
    monkeypatch.setenv('KOFBR_ML_TRAIN_DIR', str(ml_folder(tmp_path)))
    monkeypatch.setattr(apontamentos_ml, 'retrain_soon', lambda delay=20.0: None)
    apontamentos_ml._state.update(model=None, info=None, fingerprint=None, training=False)
    return apontamentos_ml


def test_ml_cleaning_and_split_without_leakage(client, ml):
    info = client.post('/api/workspace/ml/train').json
    c = info['cleaning']
    assert (c['faltantes'], c['sem_conteudo'], c['rotulo_sem_informacao']) == (2, 1, 1)
    # A linha no plural vira duplicada depois que os rótulos são unidos.
    assert c['duplicadas'] == 2 and c['conflitos'] == 1
    assert c['rotulos_antes'] - c['rotulos_depois'] == 1  # "Sensores danificados" unido a "Sensor danificado"
    assert c['exemplos'] == 75 and info['classes'] == 3
    assert info['train_size'] + info['test_size'] == 75 and info['test_size'] == 15
    raw, _, counts = ml.read_folder()
    from workflow import service
    examples, _, _ = ml.prepare(raw, [], service.catalog(), counts)
    texts = [e['text'] for e in examples]
    assert len(texts) == len(set(texts))  # um exemplo por relato: nada aparece no treino e no teste
    assert info['accuracy'] >= 80


def test_ml_fill_review_and_learn(client, ml):
    rows = [['Data', 'Equipamento', 'Observações']]
    rows += [['01/01/2026', 'ENCHEDORA', 'sensor com defeito na saida 30008216805'],
             ['01/01/2026', 'ENCHEDORA', 'sensor com defeito na saida 30008216999'],
             ['01/01/2026', 'ENCHEDORA', 'carrossel travado na entrada']]
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    data = io.BytesIO()
    workbook.save(data)
    result = client.post('/api/workspace/ml/fill', data={'file': (io.BytesIO(data.getvalue()), 'dia.xlsx')}).json
    assert result['rows'] == 3 and result['unique'] == 2 and result['preview'][0]['classe'] == 'FALHA DE SENSOR'
    items = client.get(f"/api/workspace/ml/runs/{result['id']}/items").json['items']
    carrossel = next(i for i in items if 'carrossel' in i['key'])
    sensor = next(i for i in items if 'sensor' in i['key'])
    assert sensor['n_linhas'] == 2  # mesma frase com O.S. diferentes = um relato para revisar
    review = client.post(f"/api/workspace/ml/runs/{result['id']}/review", json={'items': [
        {'key': carrossel['key'], 'acao': 'corrigir', 'classe': 'falha de carrossel', 'detalhe': 'Carrossel travado'},
        {'key': sensor['key'], 'acao': 'confirmar'}]}).json
    assert review['saved'] == 2 and review['run']['reviewed'] == 2
    # A planilha baixada já sai com a correção.
    sheet = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/workspace/ml/fill/{result['id']}").data)).active
    assert [c.value for c in sheet[4]][3:] == ['FALHA DE CARROSSEL', 'Carrossel travado', None, 'Analista']
    # O modelo retreina com as revisões e o mesmo relato já chega corrigido.
    info = client.post('/api/workspace/ml/train').json
    assert info['analyst_examples'] == 2 and 'FALHA DE CARROSSEL' in client.get('/api/workspace/ml/classes').json['classes']
    again = client.post('/api/workspace/ml/fill', data={'file': (io.BytesIO(data.getvalue()), 'dia2.xlsx')}).json
    assert again['preview'][2]['classe'] == 'FALHA DE CARROSSEL' and again['preview'][2]['faixa'] == 'Analista'
    # Revisar de novo o mesmo relato substitui a revisão anterior.
    client.post(f"/api/workspace/ml/runs/{again['id']}/review", json={'items': [
        {'key': carrossel['key'], 'acao': 'corrigir', 'classe': 'FALHA DE ESTRELA', 'detalhe': ''}]})
    from workflow import store
    examples = store.ml_examples()
    assert len(examples) == 2 and {e['classe'] for e in examples} == {'FALHA DE ESTRELA', 'FALHA DE SENSOR'}


# ---------- Fluxo único: upload → predição → validação → agrupamento → dashboard ----------
def sap_sheet(rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Apontamentos de Falha KOFBR'
    sheet.append(['Centro', 'Data Inicio Real', 'Linha', 'Tipo de Parada', 'Intervalo', 'chave do parada',
                  'Observações', 'Minutos de paradas'])
    for row in rows:
        sheet.append(row)
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def run_pipeline(client, data, name='completa.xlsx'):
    import time
    job = client.post('/api/workspace/pipeline', data={'file': (io.BytesIO(data), name)}).json
    for _ in range(300):
        status = client.get(f"/api/workspace/pipeline/jobs/{job['id']}").json
        if status['done']:
            return status
        time.sleep(0.1)
    raise AssertionError('processamento não terminou')


def test_pipeline_end_to_end(client, ml):
    rows = [
        ['B0AC', '02/01/2026', 'LINHA001', 'P.EQ.LINHA', '13:00-14:00', 'ENCHEDORA', 'sensor com defeito na saida 300082168051', 60],
        ['B0AC', '02/01/2026', 'LINHA001', 'P.EQ.LINHA', '14:00-15:00', 'ENCHEDORA', 'sensor com defeito na saida 300082168051', 30],
        ['B0AC', '03/01/2026', 'LINHA002', 'P.EQ.LINHA', '08:00-09:00', 'ROTULADORA', 'carrossel travado na entrada', 20],
        ['B0AC', '03/01/2026', '', 'P.EQ.LINHA', '09:00-10:00', '', 'rolamento com defeito', 10],  # inválida
    ]
    data = sap_sheet(rows)
    status = run_pipeline(client, data)
    assert status['error'] == '' and status['analysis_id']
    analysis_id = status['analysis_id']
    info = client.get(f'/api/workspace/pipeline/{analysis_id}').json
    assert info['rows'] == 3 and info['skipped'] == [{'linha': 5, 'motivo': 'linha e equipamento são obrigatórios'}]
    assert info['ready'] is False and info['step'] == 3
    # Dashboard ainda bloqueado: nada validado.
    assert client.get(f'/api/workspace/analytics?ids={analysis_id}').status_code == 400

    page = client.get(f'/api/workspace/pipeline/{analysis_id}/items').json
    assert page['total'] == 2 and page['progress']['linhas'] == 3  # mesma frase com O.S. diferente = 1 relato
    sensor = next(i for i in page['items'] if 'sensor' in i['relato_norm'])
    carrossel = next(i for i in page['items'] if 'carrossel' in i['relato_norm'])
    assert sensor['classe'] == 'FALHA DE SENSOR' and sensor['n_linhas'] == 2

    saved = client.post(f'/api/workspace/pipeline/{analysis_id}/validate', json={'items': [
        {'key': sensor['relato_norm'], 'resultado': 'correta'},
        {'key': carrossel['relato_norm'], 'resultado': 'errada', 'classe': 'falha de carrossel', 'detalhe': 'Carrossel travado'}]}).json
    assert saved['saved'] == 2 and saved['progress']['validados'] == 2 and saved['progress']['percent_linhas'] == 100
    # Passo 4: agrupado por falha validada (2 linhas do SAP da mesma O.S. = 1 falha real).
    summary = {r['classe']: r for r in client.get(f'/api/workspace/pipeline/{analysis_id}/summary').json['failures']}
    assert summary['FALHA DE SENSOR']['linhas'] == 2 and summary['FALHA DE SENSOR']['falhas_reais'] == 1
    assert summary['FALHA DE CARROSSEL']['linhas'] == 1 and summary['FALHA DE CARROSSEL']['linhas_validadas'] == 1
    # Passo 5: o Dashboard existente conta pela classe validada.
    analytics = client.get(f'/api/workspace/analytics?ids={analysis_id}').json
    assert {f['key']: f['lines'] for f in analytics['failures']} == {'FALHA DE SENSOR': 2, 'FALHA DE CARROSSEL': 1}
    listing = next(a for a in client.get('/api/workspace/analyses').json if a['id'] == analysis_id)
    assert listing['mode'] == 'ml' and listing['ready'] and listing['step'] == 5
    # Validações viram exemplos de treino.
    from workflow import store
    assert {e['classe'] for e in store.ml_examples()} == {'FALHA DE SENSOR', 'FALHA DE CARROSSEL'}
    # A planilha completa sai com as previsões e as correções, na linha certa.
    sheet = openpyxl.load_workbook(io.BytesIO(client.get(f'/api/workspace/pipeline/{analysis_id}/download').data)).active
    assert [c.value for c in sheet[1]][-4:] == ['Classificação Manual Analista', 'Detalhe sugerido',
                                               'Confiança do modelo (%)', 'Faixa de confiança']
    assert sheet['G4'].value == 'carrossel travado na entrada'
    assert [c.value for c in sheet[4]][-4:] == ['FALHA DE CARROSSEL', 'Carrossel travado', None, 'Analista']
    assert sheet.cell(5, 12).value == 'Linha ignorada: linha e equipamento são obrigatórios'
    # Mesmo arquivo de novo: recusado, com o link da análise existente.
    again = run_pipeline(client, data, 'copia.xlsx')
    assert again['existing_id'] == analysis_id and again['error']


def test_pipeline_accept_high_needs_confirmation(client, ml):
    data = sap_sheet([['B0AC', '02/01/2026', 'LINHA001', 'P.EQ.LINHA', '13:00-14:00', 'ENCHEDORA',
                       f'correia com defeito na posicao {i} da maquina', 10] for i in range(3)])
    analysis_id = run_pipeline(client, data)['analysis_id']
    url = f'/api/workspace/pipeline/{analysis_id}/accept-high'
    assert client.post(url, json={}).status_code == 400
    high = client.get(f'/api/workspace/pipeline/{analysis_id}/high').json
    if high['total']:
        result = client.post(url, json={'confirm': True}).json
        assert result['progress']['validados'] == high['total']


def test_kanban_groups_with_ml_shows_certainty_and_learns(client, ml):
    ml.train(force=True)
    rows = [['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'sensor com defeito na posicao 3 da maquina', 5],
            ['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'sensor com defeito na posicao 9 da maquina', 4],
            # Mesma falha, outro contexto: nunca no mesmo card.
            ['U1', '30/09/2026', 'L2', 'Enchedora', 'Mecânica', 'sensor com defeito na posicao 3 da maquina', 2],
            # Frase parecida, outra causa: outro card.
            ['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'rolamento com defeito na posicao 3 da maquina', 3]]
    board = upload(client, rows).json
    by_line = {}
    for card in board['cards']:
        col = next(c for c in board['columns'] if c['id'] == card['column_id'])
        by_line.setdefault(col['line'], []).append(card)
    l1 = {c['class']: c for c in by_line['L1']}
    assert set(l1) == {'FALHA DE SENSOR', 'FALHA DE ROLAMENTO'}
    sensor = l1['FALHA DE SENSOR']
    # Padrão: classe do catálogo; o ML só dá a % de acerto de cada card (todo card tem %).
    assert sensor['count'] == 2 and sensor['grouped_by'] == 'ml' and sensor['source'] == 'componente'
    assert 0 < sensor['ml_pct'] <= 100 and '%' in sensor['reason']
    assert all(c['ml_pct'] is not None and c['confidence'] in ('alta', 'media', 'baixa') for c in board['cards'])
    assert [c['count'] for c in by_line['L2']] == [1]
    assert board['use_ml'] is False
    # Botão "Usar machine learning": o ML refaz a classificação; desligar volta ao catálogo.
    on = action(client, board, 'use_ml', on=True).json
    assert on['use_ml'] is True and on['revision'] == board['revision'] + 1
    assert next(c for c in on['cards'] if c['id'] == sensor['id'])['source'] == 'ml'
    off = action(client, on, 'use_ml', on=False).json
    assert off['use_ml'] is False
    assert {c['id']: c['class'] for c in off['cards']} == {c['id']: c['class'] for c in board['cards']}
    board = off

    # Validar ensina o modelo: o relato volta, em outra planilha, com 100% (já validado por pessoa).
    validate_all(client, board)
    assert any(e['relato_norm'] == 'rolamento com defeito na posicao 3 da maquina' and e['classe'] == 'FALHA DE ROLAMENTO'
               for e in store.ml_examples())
    ml.train()
    again = upload(client, [['U9', '01/10/2026', 'L7', 'Rotuladora', 'Mecânica',
                             'Rolamento com defeito na posição 3 da máquina', 1]], 'outra.xlsx').json
    assert again['cards'][0]['class'] == 'FALHA DE ROLAMENTO' and again['cards'][0]['ml_pct'] == 100


def test_drag_between_cards_keeps_position(client):
    rows = [['U1', '30/09/2026', 'L1', 'Bomba', 'Mecânica', t, 4] for t in ('vazamento no selo', 'motor queimado', 'correia rompida')]
    board = upload(client, rows).json
    ids = [c['id'] for c in board['cards']]
    moved = action(client, board, 'reorder', card_id=ids[2], before_id=ids[0])
    assert moved.status_code == 200, moved.json
    assert [c['id'] for c in moved.json['cards']] == [ids[2], ids[0], ids[1]]
    # Trocar de falha soltando depois de um card: classe nova e posição escolhida.
    board = action(client, moved.json, 'set_class', card_id=ids[0], failure_class='FALHA DE MOTOR', after_id=ids[1]).json
    assert [c['id'] for c in board['cards']] == [ids[2], ids[1], ids[0]]
    assert board['cards'][2]['class'] == 'FALHA DE MOTOR'
    assert action(client, board, 'reorder', card_id=ids[0]).status_code == 400



def test_uncertain_goes_to_triage_old_board_gets_ml_and_fix_validates(client, ml, tmp_path, monkeypatch):
    import os
    folder = os.environ['KOFBR_ML_TRAIN_DIR']
    # Análise importada sem modelo (como as antigas): nada de ML nos registros.
    monkeypatch.setenv('KOFBR_ML_TRAIN_DIR', str(tmp_path / 'nao-existe'))
    rows = [['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'sensor com defeito na posicao 4 da maquina', 5],
            ['U1', '30/09/2026', 'L1', 'Enchedora', 'Mecânica', 'operador ajustou parametro do lote azul', 3]]
    board = upload(client, rows).json
    assert all(c['ml_pct'] is None for c in board['cards'])
    # Com o modelo disponível, abrir o quadro aplica o ML aos cards abertos.
    monkeypatch.setenv('KOFBR_ML_TRAIN_DIR', folder)
    fresh = client.get(f'/api/workspace/analyses/{board["id"]}').json
    assert fresh['revision'] == board['revision'] + 1
    by_text = {c['sample']: c for c in fresh['cards']}
    sensor = by_text['sensor com defeito na posicao 4 da maquina']
    assert sensor['class'] == 'FALHA DE SENSOR' and sensor['ml_pct'] >= 50 and 'sensor' in sensor['why']
    # Sem classe confiável (ML abaixo de 50% e nenhum termo do catálogo): vai para "A classificar",
    # não inventa falha, e mostra a sugestão do ML.
    rec = {'id': 1, 'observacao_raw': 'operador ajustou parametro do lote azul', 'equipamento': 'Enchedora',
           'ml_class': 'FALHA DE SENSOR', 'ml_conf': 31.0, 'ml_words': ['parametro']}
    card = {'id': 'c', 'record_ids': [1]}
    info = service.card_info(card, service.details_for([rec], service.catalog(), {}), {1: rec})
    assert info['label'] == 'SEM MODO DE FALHA IDENTIFICADO' and info['confidence'] == 'baixa'
    assert info['suggestion'] == 'FALHA DE SENSOR' and '31%' in info['why']
    other = by_text['operador ajustou parametro do lote azul']
    # Corrigir = validar: o card movido pelo analista já fica validado.
    fixed = action(client, fresh, 'set_class', card_id=other['id'], failure_class='FALHA DE AJUSTE').json
    moved = next(c for c in fixed['cards'] if c['id'] == other['id'])
    assert moved['class'] == 'FALHA DE AJUSTE' and moved['validated'] and moved['why'] == 'definida por você'
