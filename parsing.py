"""Le planilhas SAP e devolve apontamentos prontos para persistência."""
from datetime import datetime
import re
import unicodedata

import openpyxl


# Chaves normalizadas (maiúsculas, sem acento) para o nome interno usado pela API.
COLUMN_ALIASES = {
    "CENTRO": "centro",
    "DATA INICIO REAL": "data_inicio",
    "DATA DE INICIO REAL": "data_inicio",
    "LINHA": "linha",
    "TIPO DE PARADA": "tipo_parada",
    "MATERIAL": "material",
    "ORDEM": "ordem",
    "DESCRICAO DO MATERIAL": "descricao_material",
    "TURNO": "turno",
    "INTERVALO": "intervalo",
    "CHAVE DO PARADA": "equipamento",
    "CHAVE DA PARADA": "equipamento",
    "CHAVE DE PARADA": "equipamento",
    "SUBCHAVE DE PARADA": "subchave_parada",
    "CHAVE 1 DE PARADA": "sistema",
    "OBSERVACOES": "observacao_raw",
    "OBSERVACAO": "observacao_raw",
    "PROS. EFI. PERDID.": "pts_efic_perdidos",
    "PROS EFI PERDID": "pts_efic_perdidos",
    "PONTOS EFICIENCIA PERDIDOS": "pts_efic_perdidos",
    "PTOS. ACUMILADOS": "pts_acumulados",  # grafia presente em exportações SAP
    "PTOS. ACUMULADOS": "pts_acumulados",
    "PONTOS ACUMULADOS": "pts_acumulados",
    "CAIXAS PRODUZIDAS": "caixas_produzidas",
    "TOTAL MINUTOS": "total_minutos",
    "MINUTOS DE PARADAS": "minutos_parada",
    "MINUTOS DE PARADA": "minutos_parada",
    "OPERADOR": "operador",
    "USUARIO": "operador",
    "MATRICULA": "operador",
    "APONTADO POR": "operador",
}


def _norm_header(value):
    """Remove acentos e espaços extras de cabeçalhos vindos do Excel."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip().upper())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text)


def _to_iso(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value) if value is not None else None


def _to_float(value):
    if isinstance(value, str):
        # Excel normalmente já fornece número; este fallback cobre "1.234,50".
        value = value.strip().replace(".", "").replace(",", ".") if "," in value else value
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_sap_spreadsheet(file_path: str, sheet_name: str = None):
    """Retorna ``(records, warnings)`` a partir da aba SAP apropriada."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            candidate = next((s for s in wb.sheetnames if "APONTAMENTO" in _norm_header(s)), None)
            ws = wb[candidate] if candidate else wb[wb.sheetnames[0]]

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return [], ["A planilha não contém cabeçalho nem registros"]

        col_map = {
            index: COLUMN_ALIASES[normalized]
            for index, value in enumerate(header)
            if (normalized := _norm_header(value)) in COLUMN_ALIASES
        }
        found = set(col_map.values())
        required = {"linha", "equipamento", "observacao_raw", "minutos_parada"}
        warnings = [
            f"Coluna esperada não encontrada no arquivo: '{field}'"
            for field in sorted(required - found)
        ]

        records = []
        text_fields = (
            "centro", "linha", "tipo_parada", "material", "ordem", "descricao_material",
            "turno", "intervalo", "equipamento", "subchave_parada", "sistema",
            "observacao_raw", "operador",
        )
        numeric_fields = (
            "pts_efic_perdidos", "pts_acumulados", "caixas_produzidas",
            "total_minutos", "minutos_parada",
        )
        for row in rows_iter:
            if row is None or all(value is None for value in row):
                continue
            record = {field: None for field in set(COLUMN_ALIASES.values())}
            for index, internal in col_map.items():
                if index < len(row):
                    record[internal] = row[index]
            record["data_inicio"] = _to_iso(record.get("data_inicio"))
            for field in numeric_fields:
                record[field] = _to_float(record.get(field))
            for field in text_fields:
                if record.get(field) is not None:
                    record[field] = str(record[field]).strip()
            records.append(record)
        return records, warnings
    finally:
        wb.close()
