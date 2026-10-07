"""Exportação da análise validada no quadro.

- Planilha completa com Classificação (`full_from_analysis`): a planilha ORIGINAL importada, intacta,
  com a coluna "Classificação" que saiu do quadro (validada pelo analista).
- `sheet_examples`: exemplos (relato, classe) salvos antes pela tela "Classificar planilha" (removida);
  continuam valendo no treino do modelo.
"""
import io
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from . import classify as fc

RED = 'C8102E'


# ----------------------------------------------------------------------------- leitura

def _plain(text):
    text = re.sub(r'\s+', ' ', str(text or '')).strip().upper()
    return unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode()


def _header(row):
    cols = {}
    for i, value in enumerate(row or ()):
        h = re.sub(r'\s+', ' ', _plain(value)).strip()
        if not h:
            continue
        if 'text' not in cols and h.startswith('OBSERVA'):
            cols['text'] = i
        elif 'label' not in cols and h.startswith('CLASSIFICA'):
            cols['label'] = i
        elif 'machine' not in cols and (h.startswith('CHAVE DO PARADA') or h.startswith('CHAVE DA PARADA') or h == 'EQUIPAMENTO'):
            cols['machine'] = i
        elif 'line' not in cols and h == 'LINHA':
            cols['line'] = i
        elif 'unit' not in cols and h in ('CENTRO', 'UNIDADE'):
            cols['unit'] = i
        elif 'minutes' not in cols and h.startswith('MINUTOS'):
            cols['minutes'] = i
    return cols


# ----------------------------------------------------------------------------- planilha completa

FULL_COLUMN = 'Classificação'
FULL_FILLS = {'validada': 'E7F6EF', 'alta': 'E7F6EF', 'media': 'FFF6DC', 'baixa': 'FFF0F2'}
FULL_NOTE = ('Classificação padronizada (FALHA DE <componente>).\n'
             'Verde: validada no quadro ou com confiança alta.\n'
             'Amarelo: sugerida pela ML ou pelo catálogo, vale conferir.\n'
             'Vermelho: sem modo de falha identificado, análise manual.')


def write_full(data, filename, marks):
    """A planilha ORIGINAL, intacta, com a coluna Classificação preenchida.

    marks = {aba: {linha: (classe, tom)}}. A coluna entra depois da última coluna do cabeçalho, então
    nenhuma coluna existente muda de letra (fórmulas, filtros e formatação continuam valendo). Se a aba
    já tem uma coluna Classificação vazia, ela é preenchida no lugar."""
    import openpyxl
    from copy import copy
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    ext = Path(filename).suffix.lower()
    wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=ext == '.xlsm')
    for title, rows in marks.items():
        ws = wb[title] if title in wb.sheetnames else wb.worksheets[0]
        header_row, cols = 1, {}
        for r in range(1, 31):
            found = _header([c.value for c in ws[r]])
            if 'text' in found:
                header_row, cols = r, found
                break
        heads = [c.value for c in ws[header_row]]
        last = max((i + 1 for i, v in enumerate(heads) if v not in (None, '')), default=ws.max_column)
        sample = list(rows)[:300]
        if 'label' in cols and not any(ws.cell(n, cols['label'] + 1).value not in (None, '') for n in sample):
            col = cols['label'] + 1
        else:
            col = last + 1
            head = ws.cell(header_row, col)
            head.value = FULL_COLUMN
            model = ws.cell(header_row, cols.get('text', 0) + 1)
            for attr in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                setattr(head, attr, copy(getattr(model, attr)))
        ws.cell(header_row, col).comment = Comment(FULL_NOTE, 'Gargalo')
        for number, (label, tone) in rows.items():
            cell = ws.cell(number, col)
            cell.value = label
            if tone in FULL_FILLS:
                cell.fill = PatternFill('solid', fgColor=FULL_FILLS[tone])
        ws.column_dimensions[get_column_letter(col)].width = 34
        if ws.auto_filter.ref:
            start = ws.auto_filter.ref.split(':')[0]
            ws.auto_filter.ref = f'{start}:{get_column_letter(max(col, last))}{ws.max_row}'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), ('.xlsm' if ext == '.xlsm' else '.xlsx')


def _tone(label, confidence):
    if label in fc.UNCLASSIFIED:
        return 'baixa'
    return confidence or 'media'


def full_from_analysis(doc, source):
    """Planilha completa da análise: cada linha do SAP com a classe do card em que ela ficou no quadro."""
    from . import service
    board = service.board(doc)
    by_record = {}
    for card in board['cards']:
        # "Sem modo de falha" fica vermelho mesmo validado: é o que pede análise manual.
        tone = 'baixa' if card['class'] in fc.UNCLASSIFIED else 'validada' if card['validated'] else _tone(card['class'], card['confidence'])
        for rid in card['record_ids']:
            by_record[rid] = (card['class'], tone)
    stem = Path(doc.get('filename') or doc.get('name') or 'analise').stem
    if source and not str(doc.get('filename') or '').lower().endswith('.csv'):
        marks = defaultdict(dict)
        for r in doc['records']:
            if r['id'] in by_record and r.get('source_row'):
                marks[r.get('source_sheet') or ''][r['source_row']] = by_record[r['id']]
        data, ext = write_full(source, doc.get('filename') or 'x.xlsx', marks)
        return data, f'{stem} - com classificação{ext}'
    # CSV ou importação antiga sem o arquivo guardado: remonta a planilha com as células originais.
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Apontamentos'
    records = sorted(doc['records'], key=lambda r: (r.get('source_sheet') or '', r.get('source_row') or 0))
    heads = [k.split(' · ', 1)[-1] for k in (records[0].get('original') or {})] if records else []
    ws.append(heads + [FULL_COLUMN])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=RED)
    for r in records:
        label, tone = by_record.get(r['id'], ('', ''))
        ws.append(list((r.get('original') or {}).values()) + [label])
        if tone in FULL_FILLS:
            ws.cell(ws.max_row, len(heads) + 1).fill = PatternFill('solid', fgColor=FULL_FILLS[tone])
    ws.freeze_panes = 'A2'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f'{stem} - com classificação.xlsx'


# ----------------------------------------------------------------------------- agrupamento

def sheet_examples():
    from . import store
    return [(v['text'], v['label']) for v in (store.get_setting('sheet_examples') or {}).values()]
