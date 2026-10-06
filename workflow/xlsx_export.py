"""Planilha resumida no mesmo layout dos apontamentos de falha KOFBR (SAP).

Cada linha da exportação é uma falha consolidada (card validado) em um dia:
os apontamentos hora a hora da mesma falha viram uma única linha, com os
minutos somados e o intervalo do primeiro ao último horário. As colunas, a
ordem, as fontes e os destaques seguem a planilha original importada.
"""
import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from parsing import _norm_header

# Layout padrão da planilha "Apontamentos de Falha KOFBR".
DEFAULT_LAYOUT = [
    ('Centro', 'centro'), ('Data Inicio Real', 'data_inicio'), ('Linha', 'linha'), ('Tipo de Parada', 'tipo_parada'),
    ('Material', 'material'), ('Ordem', 'ordem'), ('Descrição do Material', 'descricao_material'), ('Turno', 'turno'),
    ('Intervalo', 'intervalo'), ('chave do parada', 'equipamento'), ('Subchave de Parada', 'subchave_parada'),
    ('chave 1 de parada', 'sistema'), ('Observações', 'observacao_raw'), ('Pros. Efi. Perdid.', 'pts_efic_perdidos'),
    ('Ptos. Acumilados', 'pts_acumulados'), ('Caixas Produzidas', 'caixas_produzidas'), ('Total minutos', 'total_minutos'),
    ('Minutos de paradas', 'minutos_parada'),
]
WIDTHS = {'centro': 8.6, 'data_inicio': 15.7, 'linha': 10, 'tipo_parada': 14.9, 'material': 9.7, 'ordem': 10,
          'descricao_material': 41.9, 'turno': 7.6, 'intervalo': 12, 'equipamento': 37.1, 'subchave_parada': 22.9,
          'sistema': 37.0, 'observacao_raw': 92.9, 'pts_efic_perdidos': 18, 'pts_acumulados': 17.1,
          'caixas_produzidas': 20.1, 'total_minutos': 14.0, 'minutos_parada': 24.4}
NUMBER_FORMATS = {'pts_efic_perdidos': '0.0000', 'pts_acumulados': '0.0000', 'caixas_produzidas': '0',
                  'total_minutos': '0.000', 'minutos_parada': '0.000'}
SUM_FIELDS = {'pts_efic_perdidos', 'caixas_produzidas', 'total_minutos', 'minutos_parada'}
MAX_FIELDS = {'pts_acumulados'}

YELLOW = PatternFill('solid', fgColor='FFFF00')
RED = PatternFill('solid', fgColor='C00000')
GREY = PatternFill('solid', fgColor='D9D9D9')
THIN = Side(style='thin', color='000000')
HEAD_FONT = Font(name='Aptos Narrow', size=11)
RED_HEAD_FONT = Font(name='Verdana', size=10, color='FFFFFF')
CELL_FONT = Font(name='Arial', size=10)
CENTER = Alignment(horizontal='center', vertical='center')
INTERVAL = re.compile(r'^\s*(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})\s*$')


def _layout(records):
    """Cabeçalhos na ordem da planilha importada, com o campo interno de cada um."""
    from workflow.importer import ALIASES
    for rec in records:
        original = rec.get('original') or {}
        if original:
            cols = []
            for key in sorted(original, key=lambda k: int(k.split(' · ', 1)[0]) if k.split(' · ', 1)[0].isdigit() else 0):
                header = key.split(' · ', 1)[-1]
                cols.append((header, ALIASES.get(_norm_header(header)), key))
            return cols
    return [(h, f, None) for h, f in DEFAULT_LAYOUT]


def _number(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(',', '.'))
    except ValueError:
        return None


def _join(values, limit=4):
    seen = []
    for v in values:
        v = '' if v is None else str(v).strip()
        if v and v not in seen:
            seen.append(v)
    if not seen:
        return ''
    if len(seen) == 1:
        return seen[0]
    return ' / '.join(seen[:limit]) + (' …' if len(seen) > limit else '')


def _interval(values):
    spans = [INTERVAL.match(str(v or '')) for v in values if str(v or '').strip()]
    if spans and all(spans):
        to_min = lambda t: int(t.split(':')[0]) * 60 + int(t.split(':')[1])
        start = min((m.group(1) for m in spans), key=to_min)
        end = max((m.group(2) for m in spans), key=to_min)
        return f'{start.zfill(5)}-{end.zfill(5)}'
    return _join(values)


def _turno(values):
    joined = _join(values)
    try:
        return int(float(joined)) if joined and ' / ' not in joined else joined
    except ValueError:
        return joined


def _value(rows, header, field, key):
    """Valor consolidado de uma coluna para um grupo de apontamentos."""
    first = rows[0]
    if field == 'data_inicio':
        return datetime.combine(date.fromisoformat(first['data_inicio']), datetime.min.time())
    if field in ('centro', 'linha', 'equipamento'):
        return first[field]  # coluna validada no quadro
    # Valores brutos da planilha original (preserva zeros e o texto exato).
    raw = [(r.get('original') or {}).get(key, '') if key else r.get(field) for r in rows]
    if field == 'observacao_raw':
        texts = [str(v).strip() for v in raw if str(v or '').strip()]
        return max(texts, key=texts.count) if texts else first['failure_mode']
    if field in SUM_FIELDS or field in MAX_FIELDS:
        nums = [n for n in (_number(v) for v in raw) if n is not None]
        if not nums:
            return None
        return round(max(nums) if field in MAX_FIELDS else sum(nums), 4)
    if field == 'intervalo':
        return _interval(raw)
    if field == 'turno':
        return _turno(raw)
    return _join(raw)


def _head(cell, field, text):
    cell.value = text
    cell.alignment = CENTER
    if field in ('observacao_raw', 'minutos_parada'):
        cell.fill, cell.font = RED, RED_HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    elif field == 'equipamento':
        cell.fill, cell.font = YELLOW, HEAD_FONT
    else:
        cell.font = HEAD_FONT


DOC_RED = PatternFill('solid', fgColor='C8102E')


DOC_HEADS = ['', None, 'T (min)', 'Q', 'MTTR', 'T.T.A', '%', '80%', 'Categoria']
DOC_WIDTHS = [6, 46, 11, 8, 9, 11, 8, 11, 20]


def _doc_block(sheet, top, first_col, items, min_rows):
    """Tabela no formato da apresentação: T, Q, MTTR, T.T.A (tempo acumulado),
    % acumulado e a coluna 80% (tempo só dos itens dentro dos 80% do Pareto)."""
    for col, h in enumerate(DOC_HEADS, 1):
        cell = sheet.cell(row=top, column=col, value=first_col if h is None else h)
        cell.fill, cell.font = DOC_RED, Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='left' if col == 2 else 'center', vertical='center')
    total_min = sum(m['minutes'] for _, m in items) or 1
    acc = 0.0
    total = max(min_rows, len(items))
    for i in range(total):
        row = top + 1 + i
        sheet.cell(row=row, column=1, value=i + 1).alignment = CENTER
        if i < len(items):
            name, m = items[i]
            acc += m['minutes']
            pct = acc / total_min
            values = [name, m['minutes'], m['count'], m['mttr'], round(acc, 1), pct,
                      m['minutes'] if pct <= 0.8 + 1e-9 else None, m['category'].upper()]
            for col, value in enumerate(values, 2):
                cell = sheet.cell(row=row, column=col, value=value)
                cell.alignment = Alignment(horizontal='left' if col == 2 else 'center')
                cell.number_format = {3: '0.0', 5: '0.0', 6: '0.0', 7: '0%', 8: '0.0'}.get(col, 'General')
        for col in range(1, len(DOC_HEADS) + 1):
            c = sheet.cell(row=row, column=col)
            c.font = Font(name='Arial', size=10, color='595959' if col == 1 else '000000')
            c.border = Border(bottom=Side(style='thin', color='E5E1DE'))
    return top + 1 + total


def _doc_table(sheet, first_col, items, note):
    for col, w in enumerate(DOC_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = w
    end = _doc_block(sheet, 1, first_col, items, 40)
    sheet.freeze_panes = 'A2'
    sheet.cell(row=end + 1, column=2, value=note).font = Font(name='Arial', size=9, italic=True, color='595959')


def _drill_sheet(sheet, machines, records, mode, note):
    """Desdobramento da apresentação: para cada máquina dos 80%, os modos de falha dela."""
    from .service import _failures, _counter
    for col, w in enumerate(DOC_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = w
    by_machine = defaultdict(list)
    for r in records:
        by_machine['\x1f'.join((r['centro'], r['linha'], r['equipamento']))].append(r)
    top = 1
    for name, m in machines:
        if top > 1 and m['cumulative'] > 80 + 1e-9 and m is not machines[0][1]:
            break
        items = [(f['name'], f) for f in _failures(by_machine.get(m['key'], []), _counter(mode))['failures']]
        cell = sheet.cell(row=top, column=2, value=f"{name} · {m['minutes']:.1f} min · Q {m['count']} · {m['category'].upper()}")
        cell.font = Font(name='Arial', size=11, bold=True)
        top = _doc_block(sheet, top + 1, name, items, 0) + 2
    sheet.cell(row=top, column=2, value=note).font = Font(name='Arial', size=9, italic=True, color='595959')


def build_summary(records, data, filenames, start, end, filters):
    """records: apontamentos validados do filtro (já com equipamento/linha/centro do quadro)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Apontamentos Resumidos KOFBR'
    layout = _layout(records)
    category = {m['key']: m['category'] for m in data['machines']}

    groups = defaultdict(list)
    for r in records:
        groups[(r['group_id'], r['data_inicio'])].append(r)
    rows = sorted(groups.values(), key=lambda g: (g[0]['data_inicio'], g[0]['centro'], g[0]['linha'],
                                                  g[0]['equipamento'], -sum(r['minutos_parada'] for r in g)))

    extra = [('Classe de falha', 'class'), ('Qtd. apontamentos', 'count'), ('Falhas (O.S.)', 'events'),
             ('O.S.', 'os'), ('Descrição consolidada', 'mode'), ('Classificação Jack-Knife', 'category')]
    for col, (header, field, _) in enumerate(layout, 1):
        _head(ws.cell(row=1, column=col), field, header)
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS.get(field, max(10, min(40, len(header) + 4)))
    base = len(layout)
    for i, (header, _) in enumerate(extra, 1):
        cell = ws.cell(row=1, column=base + i, value=header)
        cell.font, cell.alignment, cell.fill = HEAD_FONT, CENTER, GREY
        ws.column_dimensions[get_column_letter(base + i)].width = (34, 18, 14, 16, 48, 24)[i - 1]
    ws.row_dimensions[1].height = 15.95

    for r_index, group in enumerate(rows, 2):
        for col, (header, field, key) in enumerate(layout, 1):
            cell = ws.cell(row=r_index, column=col, value=_value(group, header, field, key))
            cell.font, cell.alignment = CELL_FONT, CENTER
            if field == 'data_inicio':
                cell.number_format = 'mm-dd-yy'
            elif field in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[field]
        first = group[0]
        machine_key = '\x1f'.join((first['centro'], first['linha'], first['equipamento']))
        klass = Counter(r.get('failure_class') or '' for r in group).most_common(1)[0][0]
        events = len({r.get('event_id') or r['id'] for r in group})
        orders = ', '.join(sorted({r['os'] for r in group if r.get('os')})[:4])
        for i, value in enumerate([klass, len(group), events, orders, first['failure_mode'], category.get(machine_key, '')], 1):
            cell = ws.cell(row=r_index, column=base + i, value=value)
            cell.font, cell.alignment = CELL_FONT, CENTER
    ws.freeze_panes = 'A2'
    for sheet_ in (ws,):
        sheet_.page_setup.orientation = 'landscape'
        sheet_.page_setup.fitToWidth, sheet_.page_setup.fitToHeight = 1, 0
        sheet_.sheet_properties.pageSetUpPr.fitToPage = True
        sheet_.print_title_rows = '1:1'
    if rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(base + len(extra))}{len(rows) + 1}'

    # Tabelas do descritivo (item 3.2): sequência numérica com no mínimo 40 linhas.
    from .service import short_line
    period = (f'{start} a {end}' if start != end else start) if start else 'todo o período'
    note = (f"Período: {period} · Arquivo(s): {', '.join(filenames) or 'todas as análises validadas'}"
            + ''.join(f' · {k}: {v}' for k, v in filters.items())
            + f" · {data['metrics']['lines']} linhas do SAP = {data['metrics']['events']} falhas (mesma O.S. = 1 falha)"
            + f" · Q contado em {'falhas' if data.get('count_mode', 'events') == 'events' else 'linhas do SAP'}"
            + f" · {len(rows)} linhas resumidas.")
    machines = [(f"{short_line(m['line'])}_{m['name']}", m) for m in data['machines']]
    failures = [(f['name'], f) for f in data.get('failures', [])]
    for title, first_col, items in (('Tabela de máquinas', 'EQUIPAMENTOS', machines),
                                    ('Tabela de falhas', 'FALHAS', failures)):
        _doc_table(wb.create_sheet(title), first_col, items, note)
    _drill_sheet(wb.create_sheet('Falhas por máquina'), machines, records, data.get('count_mode', 'events'), note)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
