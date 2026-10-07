"""Fluxo único (human-in-the-loop): Upload → Predição do ML → Validação humana → Agrupamento → Dashboard.

Como os dados ficam guardados (pensado para a planilha completa, ~94 mil linhas):
- radar_analyses: um documento leve por análise (mode = 'ml'), com os apontamentos sem as células
  originais. É gravado só no upload; as telas de Dashboard e Comparar o leem como qualquer análise.
- radar_ml_items: uma linha por relato único (a mesma frase com O.S. diferentes é um relato só) com a
  previsão do modelo e a validação do analista. Validar um relato vale para todas as linhas dele.
- radar_failure_summary: o Passo 4, ocorrências agrupadas por falha, recalculado a cada "Salvar".
- data/ml_saidas/<id>.xlsx: a planilha enviada com as previsões preenchidas; as validações são
  aplicadas na hora do download.

O processamento do upload (~1 a 2 min na planilha completa) roda numa thread; a tela acompanha o job.
"""
import io
import json
import re
import threading
import time
import traceback
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from . import apontamentos_ml as ml
from . import classify as fc
from . import importer, store

STEPS = ['Lendo a planilha', 'Prevendo a falha de cada relato (ML)', 'Escrevendo as previsões na planilha',
         'Gravando no banco', 'Agrupando por falha']
CORRECT, WRONG, PENDING = 'correta', 'corrigida', 'pendente'
_jobs = {}
_file_lock = threading.Lock()


def _now():
    return datetime.now().astimezone().isoformat()


# ----------------------------------------------------------------------------- Passos 1 e 2

def start(data, filename):
    """Passo 1: recebe o arquivo e dispara o processamento em segundo plano."""
    filename = Path(str(filename).replace('\\', '/')).name
    if not filename.lower().endswith(ml.EXTENSIONS):
        raise ValueError('Envie a planilha completa em .xlsx ou .xlsm.')
    if not data:
        raise ValueError('O arquivo está vazio.')
    job = {'id': uuid4().hex, 'filename': filename, 'step': 0, 'steps': STEPS, 'done': False, 'error': '',
           'existing_id': '', 'analysis_id': '', 'started_at': time.time(), 'seconds': 0, 'detail': ''}
    _jobs[job['id']] = job
    threading.Thread(target=_run, args=(job, data, filename), daemon=True).start()
    return job


def job_status(job_id):
    job = _jobs.get(job_id)
    if not job:
        raise ValueError('Processamento não encontrado. Envie a planilha novamente.')
    return {**job, 'seconds': round((time.time() - job['started_at']) if not job['done'] else job['seconds'], 1)}


def _run(job, data, filename):
    try:
        process(data, filename, job)
    except Exception as exc:  # o erro aparece na tela, no passo em que parou
        job['error'] = str(exc) if isinstance(exc, ValueError) else 'Falha inesperada no processamento.'
        if not isinstance(exc, ValueError):
            traceback.print_exc()
    finally:
        job['seconds'] = round(time.time() - job['started_at'], 1)
        job['done'] = True


def process(data, filename, job=None):
    """Passo 2 completo. Devolve o id da análise criada."""
    job = job if job is not None else {}

    def step(n, detail=''):
        job['step'], job['detail'] = n, detail

    step(0)
    skipped = []
    records, warnings = importer.parse_file(data, filename, strict=False, keep_original=False, skipped=skipped)
    fingerprint = importer.content_hash(records)
    with store.connection() as conn:
        previous = conn.execute('SELECT id FROM radar_analyses WHERE content_hash = ?', (fingerprint,)).fetchone()
    if previous:
        job['existing_id'] = previous['id']
        raise ValueError('Esta planilha já foi enviada. Abra a análise existente para continuar a validação.')

    step(1, f'{len(records):,} linhas'.replace(',', '.'))
    for r in records:
        key = ml.normalize_text(r.get('observacao_raw'))
        r['ml_key'] = key if ml.has_content(key) else ''
    predictions = ml.predict([r['ml_key'] for r in records if r['ml_key']])
    items = _items(records, predictions)

    step(2)
    analysis_id = str(uuid4())
    _write_workbook(data, filename, analysis_id, records, predictions, skipped)

    step(3)
    model = ml.status()
    doc = {'id': analysis_id, 'mode': 'ml', 'revision': 1, 'name': filename.rsplit('.', 1)[0], 'filename': filename,
           'created_at': _now(), 'updated_at': _now(), 'records': records, 'columns': [], 'cards': [],
           'warnings': warnings, 'events': [],
           'skipped': [{'linha': n, 'motivo': why} for n, why in skipped],
           'ml': {'accuracy': model.get('accuracy'), 'margin': model.get('margin'), 'trained_at': model.get('trained_at')}}
    _, inserted = store.insert(doc, fingerprint, data)
    if not inserted:
        raise ValueError('Esta planilha já foi enviada.')
    store.insert_ml_items(analysis_id, items)

    step(4)
    summarize(analysis_id)
    job['analysis_id'] = analysis_id
    return analysis_id


def _items(records, predictions):
    """Um item por relato único, com a previsão do modelo."""
    groups = defaultdict(list)
    for r in records:
        if r['ml_key']:
            groups[r['ml_key']].append(r)
    items = []
    for key, rows in groups.items():
        p = predictions[key]
        # Relato que o analista já validou antes (memória do modelo) chega validado.
        known = p['faixa'] == ml.ANALYST
        items.append({'relato_norm': key, 'relato': str(rows[0]['observacao_raw']), 'n_linhas': len(rows),
                      'minutos': round(sum(r['minutos_parada'] for r in rows), 2), 'classe': p['classe'],
                      'detalhe': p['detalhe'], 'confianca': p['confianca'], 'faixa': p['faixa'], 'top3': p['top3'],
                      'status': CORRECT if known else PENDING, 'classe_final': p['classe'] if known else None,
                      'detalhe_final': p['detalhe'] if known else None, 'updated_at': _now() if known else None})
    return items


def _layout_path(analysis_id):
    return ml._outputs_dir() / f'{analysis_id}.layout.json'


def _write_workbook(data, filename, analysis_id, records, predictions, skipped):
    """A planilha enviada, com as previsões nas colunas novas (mesmo layout da tela Gerar Planilha)."""
    import openpyxl
    from openpyxl.styles import Font
    ext = Path(filename).suffix.lower()
    wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=ext == '.xlsm')
    sheet_name = records[0]['source_sheet'] if records else wb.worksheets[0].title
    ws = wb[sheet_name]
    head = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=1, max_row=30)]
    header, _, label_col = ml._find_header(head)
    header_row = (header or 0) + 1
    last = ws.max_column
    bold = Font(bold=True)
    if label_col is None:
        last += 1
        label_col = last
        ws.cell(header_row, label_col, ml.TARGET).font = bold
    else:
        label_col += 1
    cols = {'label': label_col, 'detail': last + 1, 'conf': last + 2, 'band': last + 3}
    for key, title in (('detail', ml.DETAIL_COLUMN), ('conf', ml.CONF_COLUMN), ('band', ml.BAND_COLUMN)):
        ws.cell(header_row, cols[key], title).font = bold
    rows_by_key = defaultdict(list)
    for r in records:
        if r['ml_key']:
            ml._paint(ws, r['source_row'], cols, predictions[r['ml_key']])
            rows_by_key[r['ml_key']].append(r['source_row'])
        else:
            ws.cell(r['source_row'], cols['band']).value = 'Sem relato para classificar'
    for number, why in skipped:
        ws.cell(number, cols['band']).value = f'Linha ignorada: {why}'
    for key, width in (('label', 34), ('detail', 40), ('band', 22)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cols[key])].width = width
    out_ext = '.xlsm' if ext == '.xlsm' else '.xlsx'
    wb.save(ml._outputs_dir() / f'{analysis_id}{out_ext}')
    _layout_path(analysis_id).write_text(json.dumps({
        'sheet': sheet_name, 'cols': cols, 'ext': out_ext, 'dirty': False,
        'filename': f'{Path(filename).stem} - classificada{out_ext}', 'rows': rows_by_key}, ensure_ascii=False),
        encoding='utf-8')


# ----------------------------------------------------------------------------- Passo 3

def items_page(analysis_id, faixa='', status='', q='', page=1):
    _document(analysis_id)
    total, items = store.ml_items_page(analysis_id, faixa=faixa, status=status, q=q, page=page)
    return {'total': total, 'page': max(page, 1), 'size': 100, 'items': items, 'progress': progress(analysis_id)}


def progress(analysis_id):
    p = store.ml_items_progress(analysis_id)
    p['percent_linhas'] = round(p['linhas_validadas'] * 100 / p['linhas'], 1) if p['linhas'] else 0
    return p


def validate(analysis_id, entries):
    """Salva 'Correta' / 'Errada + classe correta'. Vira exemplo de treino e recalcula o agrupamento."""
    doc = _document(analysis_id)
    if not isinstance(entries, list) or not entries:
        raise ValueError('Nenhuma validação para salvar.')
    entries = entries[:2000]
    current = {i['relato_norm']: i for i in store.ml_items_by_keys(analysis_id, [str((e or {}).get('key') or '') for e in entries])}
    updates, examples = [], []
    now = _now()
    for entry in entries:
        item = current.get(str(entry.get('key') or ''))
        if not item:
            continue
        result = entry.get('resultado')
        if result == 'correta':
            classe, detalhe, status = item['classe'], item['detalhe'] or '', CORRECT
        elif result == 'errada':
            classe = ml.class_label(entry.get('classe'))
            detalhe = re.sub(r'\s+', ' ', str(entry.get('detalhe') or '')).strip()[:300]
            if not classe:
                raise ValueError('Informe a classificação correta dos relatos marcados como errados.')
            status = WRONG
        elif result == 'pendente':
            classe = detalhe = None
            status = PENDING
        else:
            raise ValueError('Marque cada relato como correta ou errada.')
        updates.append({'relato_norm': item['relato_norm'], 'status': status, 'classe_final': classe,
                        'detalhe_final': detalhe, 'updated_at': now})
        if status != PENDING:
            examples.append({'relato': item['relato'], 'relato_norm': item['relato_norm'], 'classe': classe,
                             'detalhe': detalhe, 'origem': doc['filename']})
    if not updates:
        raise ValueError('Os relatos enviados não pertencem a esta análise.')
    store.update_ml_items(analysis_id, updates)
    if examples:
        store.save_ml_examples(examples)
        ml.retrain_soon()
    _mark_dirty(analysis_id)
    summary = summarize(analysis_id)
    return {'saved': len(updates), 'progress': progress(analysis_id), 'summary': summary}


def accept_high(analysis_id, confirm, keys=None):
    """Lote seguro: marca como corretas as previsões de confiança Alta ainda pendentes."""
    if confirm is not True:
        raise ValueError('Confirme que conferiu a amostra antes de validar em lote.')
    items = [i for i in store.ml_items(analysis_id) if i['status'] == PENDING and i['faixa'] == 'Alta']
    if keys:
        wanted = set(keys)
        items = [i for i in items if i['relato_norm'] in wanted]
    if not items:
        raise ValueError('Nenhum relato de confiança alta pendente.')
    return validate(analysis_id, [{'key': i['relato_norm'], 'resultado': 'correta'} for i in items])


def high_sample(analysis_id, n=8):
    items = sorted((i for i in store.ml_items(analysis_id) if i['status'] == PENDING and i['faixa'] == 'Alta'),
                   key=lambda i: -i['n_linhas'])
    step = max(1, len(items) // n)
    return {'total': len(items), 'linhas': sum(i['n_linhas'] for i in items), 'sample': items[::step][:n]}


# ----------------------------------------------------------------------------- Passo 4

def final_classes(analysis_id):
    """relato normalizado → (classe que vale, validado?)"""
    return {i['relato_norm']: (i['classe_final'] or i['classe'], i['status'] != PENDING)
            for i in store.ml_items(analysis_id)}


def summarize(analysis_id):
    """Agrupa as ocorrências por falha (validada; senão, a prevista) e grava o resultado."""
    from . import service
    doc = _document(analysis_id)
    classes = final_classes(analysis_id)
    events = service.event_map(doc)
    by_class = defaultdict(lambda: {'linhas': 0, 'eventos': set(), 'minutos': 0.0, 'validadas': 0})
    for r in doc['records']:
        cls, checked = classes.get(r.get('ml_key'), (fc.SEM_DESCRICAO, False))
        g = by_class[cls]
        g['linhas'] += 1
        g['eventos'].add(events[r['id']][0])
        g['minutos'] += r['minutos_parada']
        g['validadas'] += checked
    total = sum(g['minutos'] for g in by_class.values()) or 1
    rows, cumulative, now = [], 0.0, _now()
    for cls, g in sorted(by_class.items(), key=lambda kv: (-kv[1]['minutos'], kv[0])):
        cumulative += g['minutos']
        falhas = len(g['eventos'])
        rows.append({'classe': cls, 'linhas': g['linhas'], 'falhas_reais': falhas, 'minutos': round(g['minutos'], 2),
                     'mttr': round(g['minutos'] / falhas, 2) if falhas else 0,
                     'percentual': round(g['minutos'] * 100 / total, 2),
                     'acumulado': round(min(100.0, cumulative * 100 / total), 2),
                     'linhas_validadas': g['validadas'], 'updated_at': now})
    store.save_failure_summary(analysis_id, rows)
    return rows


def summary(analysis_id):
    doc = _document(analysis_id)
    return {'analysis': analysis_info(doc), 'failures': store.failure_summary(analysis_id)}


# ----------------------------------------------------------------------------- integração

def analysis_info(doc):
    p = progress(doc['id'])
    layout = _read_layout(doc['id']) or {}
    return {'id': doc['id'], 'name': doc['name'], 'filename': doc['filename'], 'created_at': doc['created_at'],
            'rows': len(doc['records']), 'skipped': doc.get('skipped', []), 'warnings': doc.get('warnings', []),
            'model': doc.get('ml', {}), 'progress': p, 'ready': p['validados'] > 0,
            'step': 5 if p['validados'] else 3,
            'download': layout.get('filename', '')}


def _read_layout(analysis_id):
    path = _layout_path(analysis_id)
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else None


def records_with_classes(doc):
    """Apontamentos da análise com a classe que vale (usado pelo Dashboard e pelo Comparar)."""
    from . import service
    classes = final_classes(doc['id'])
    events = service.event_map(doc)
    for r in doc['records']:
        cls, _ = classes.get(r.get('ml_key'), (fc.SEM_DESCRICAO, False))
        yield {**r, 'failure_class': cls, 'failure_mode': cls, 'group_id': r.get('ml_key') or f"__{r['id']}",
               'analysis_id': doc['id'], 'event_id': events[r['id']][0], 'os': events[r['id']][1]}


def is_ready(doc):
    return store.ml_items_progress(doc['id'])['validados'] > 0


def _mark_dirty(analysis_id):
    with _file_lock:
        path = _layout_path(analysis_id)
        if path.exists():
            layout = json.loads(path.read_text(encoding='utf-8'))
            layout['dirty'] = True
            path.write_text(json.dumps(layout, ensure_ascii=False), encoding='utf-8')


def output_file(analysis_id):
    """Planilha completa para download, com as validações aplicadas."""
    _document(analysis_id)
    with _file_lock:
        path = _layout_path(analysis_id)
        if not path.exists():
            raise ValueError('Planilha com as previsões não encontrada.')
        layout = json.loads(path.read_text(encoding='utf-8'))
        xlsx = ml._outputs_dir() / f"{analysis_id}{layout['ext']}"
        if layout.get('dirty'):
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, keep_vba=layout['ext'] == '.xlsm')
            ws = wb[layout['sheet']]
            for item in store.ml_items(analysis_id):
                if item['status'] == PENDING:
                    continue
                painted = {'classe': item['classe_final'], 'detalhe': item['detalhe_final'], 'confianca': None,
                           'faixa': ml.ANALYST}
                for row in layout['rows'].get(item['relato_norm'], []):
                    ml._paint(ws, row, layout['cols'], painted)
            wb.save(xlsx)
            layout['dirty'] = False
            path.write_text(json.dumps(layout, ensure_ascii=False), encoding='utf-8')
        return xlsx, layout['filename']


def _document(analysis_id):
    doc = store.get(analysis_id)
    if not doc or doc.get('mode') != 'ml':
        raise ValueError('Análise do fluxo único não encontrada.')
    return doc
