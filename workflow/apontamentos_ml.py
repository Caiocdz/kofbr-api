"""Preenchimento automático da coluna 'Classificação Manual Analista' (scikit-learn, 100% local).

Dados de treino
- Planilhas da pasta de treino (padrão: "planilhas modelo/plhanilha treinamento de ml", ou
  KOFBR_ML_TRAIN_DIR). Em cada aba, procura nas primeiras 30 linhas um cabeçalho "Observações..."
  (relato) e um "Classificação..." (rótulo do analista).
- Exemplos revisados pelo analista na tela "Revisar previsões" (tabela radar_ml_examples). Valem 3x,
  vencem a pasta quando o relato é o mesmo e entram SEMPRE no treino (nenhuma correção se perde);
  o teste de 20% é sorteado só entre os exemplos da pasta, que o modelo nunca viu.

Tratamento (cada etapa é contada e mostrada na tela)
- descarta linhas sem relato ou sem rótulo, relatos sem conteúdo ("xxxxxx", "nota", só a O.S.) e
  rótulos que não informam nada ("Sem detalhes", "Apontamento errado");
- une rótulos quase iguais (singular/plural, preposições, digitação);
- junta duplicadas (mesmo relato + mesmo rótulo) num só exemplo com peso = quantidade;
- relato repetido com rótulos diferentes fica com o rótulo da maioria.
Depois disso cada relato aparece uma única vez, então a divisão 80/20 nunca põe o mesmo relato no
treino e no teste (antes isso inflava a acurácia).

Dois níveis
- Classe padronizada (coluna principal): o rótulo do analista é enquadrado no catálogo de falhas do
  quadro ("FALHA DE SENSOR"...); se não enquadrar, tenta pelo relato; rótulos recorrentes (5+) que o
  catálogo não conhece viram classe própria. Modelo: TF-IDF de palavras + pedaços de palavras com SVM
  linear; confiança = softmax das pontuações (temperatura 0,15), com a acurácia de cada faixa medida
  no teste.
- Detalhe sugerido: dentro da classe prevista, o rótulo de analista do relato de treino mais parecido.

Avaliação: 80/20 com semente fixa (o modelo usado é o dos 80%), IC 95% de Wilson e, para estabilidade,
o acerto médio em 5 sorteios diferentes.
"""
import hashlib
import io
import json
import math
import os
import random
import re
import threading
import time
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from uuid import uuid4

from . import classify as fc

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TRAIN_DIR = ROOT / 'planilhas modelo' / 'plhanilha treinamento de ml'
TARGET = 'Classificação Manual Analista'
DETAIL_COLUMN = 'Detalhe sugerido'
CONF_COLUMN = 'Confiança do modelo (%)'
BAND_COLUMN = 'Faixa de confiança'
ANALYST = 'Analista'
# Faixas de confiança (probabilidade da classe escolhida). A acurácia real de cada faixa é medida no teste.
HIGH, MEDIUM = 0.5, 0.2
TEMPERATURE = 0.15
TEST_SIZE = 0.2
SEED = 42
STABILITY_SEEDS = (42, 7, 1, 3, 11)
C_GRID = (0.5, 1.0, 2.0)
ANALYST_WEIGHT = 3
MIN_OWN_CLASS = 5        # rótulo fora do catálogo vira classe própria a partir de 5 exemplos
LABEL_SIMILARITY = 0.88  # rótulos com essa semelhança (ou mais) são unidos
EXTENSIONS = ('.xlsx', '.xlsm')
JUNK_WORDS = {'xxxxxx', 'xxx', 'xx', 'x', 'nota', 'os', 'falha', 'ajuste', 'teste', 'obs', 'sem', 'nada'}
NO_INFO_LABEL = re.compile(
    r'^(sem (mais )?(detalhes?|informac\w*|descric\w*)( da falha)?|apontamento errado|nao informad\w*|'
    r'outros?|nada|indefinido)$')
LABEL_STOPWORDS = {'de', 'da', 'do', 'das', 'dos', 'na', 'no', 'nas', 'nos', 'e', 'a', 'o', 'em', 'com'}

_lock = threading.Lock()
_runs_lock = threading.Lock()
_state = {'model': None, 'info': None, 'fingerprint': None, 'training': False}
_timer = {'t': None}


def train_dir():
    return Path(os.environ.get('KOFBR_ML_TRAIN_DIR', DEFAULT_TRAIN_DIR))


def _data_dir():
    base = Path(os.environ.get('KOFBR_DATA_DIR', ROOT / 'data'))
    base.mkdir(parents=True, exist_ok=True)
    return base


def _model_path():
    return _data_dir() / 'apontamentos_ml.joblib'


def _outputs_dir():
    path = _data_dir() / 'ml_saidas'
    path.mkdir(parents=True, exist_ok=True)
    return path


def _plain(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def normalize_text(text):
    """Relato sem acentos, caixa, pontuação e números de O.S./nota (5+ dígitos)."""
    return re.sub(r'\s+', ' ', re.sub(r'\b\d{5,}\b', ' ', _plain(text))).strip()


def label_key(label):
    return _plain(label)


def has_content(norm):
    """O relato diz alguma coisa além de números e palavras de preenchimento?"""
    words = [w for w in norm.split() if not w.isdigit() and w not in JUNK_WORDS]
    return len(' '.join(words)) >= 4


def class_label(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().upper()[:120]


def _header_columns(row):
    text_col = label_col = None
    for i, value in enumerate(row):
        h = _plain(value)
        if text_col is None and h.startswith('observa'):
            text_col = i
        elif label_col is None and h.startswith('classifica'):
            label_col = i
    return text_col, label_col


def _find_header(rows):
    """(índice da linha de cabeçalho, coluna do relato, coluna do rótulo) nas primeiras 30 linhas."""
    for index, row in enumerate(rows[:30]):
        text_col, label_col = _header_columns(row or ())
        if text_col is not None:
            return index, text_col, label_col
    return None, None, None


# ----------------------------------------------------------------------------- dados de treino

def _training_files():
    folder = train_dir()
    if not folder.is_dir():
        raise ValueError(f'Pasta de treino não encontrada: {folder}')
    files = sorted(p for p in folder.iterdir() if p.suffix.lower() in EXTENSIONS and not p.name.startswith('~$'))
    if not files:
        raise ValueError(f'Nenhuma planilha .xlsx/.xlsm na pasta de treino: {folder}')
    return files


def _analyst_examples():
    try:
        from . import store
        return store.ml_examples()
    except Exception:
        return []


def _fingerprint(files, analyst):
    h = hashlib.sha1()
    for f in files:
        st = f.stat()
        h.update(f'{f.name}|{st.st_size}|{int(st.st_mtime)}'.encode())
    for ex in sorted(analyst, key=lambda e: e['relato_norm']):
        h.update(f"|{ex['relato_norm']}|{ex['classe']}|{ex['detalhe']}".encode())
    return h.hexdigest()


def read_folder():
    """Linhas brutas (relato, rótulo, arquivo) da pasta de treino + contagem de linhas lidas/faltantes."""
    import openpyxl
    rows, per_file, read, missing = [], [], 0, 0
    for path in _training_files():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        used = 0
        for ws in wb.worksheets:
            data = list(ws.iter_rows(values_only=True))
            header, text_col, label_col = _find_header(data)
            if header is None or label_col is None:
                continue
            for row in data[header + 1:]:
                if not row or all(v in (None, '') for v in row):
                    continue
                read += 1
                text = row[text_col] if text_col < len(row) else None
                label = row[label_col] if label_col < len(row) else None
                if not str(text or '').strip() or not str(label or '').strip():
                    missing += 1
                    continue
                rows.append((str(text), re.sub(r'\s+', ' ', str(label)).strip(), path.name))
                used += 1
        wb.close()
        per_file.append({'file': path.name, 'rows': used})
    return rows, per_file, {'lidas': read, 'faltantes': missing}


def _canon(key):
    words = [re.sub(r'(oes|aes|ns|s)$', '', w) for w in key.split() if w not in LABEL_STOPWORDS]
    return ' '.join(sorted(words))


def merge_labels(keys):
    """Chave do rótulo → chave do grupo de rótulos quase iguais (o mais frequente representa o grupo)."""
    freq = Counter(keys)
    by_canon, reps, merged = {}, [], {}
    for key, _ in freq.most_common():
        canon = _canon(key)
        if canon in by_canon:
            merged[key] = by_canon[canon]
            continue
        # Os testes rápidos (tamanho e letras em comum) descartam a maioria dos pares antes do cálculo caro.
        target = None
        for rep, rep_canon in reps:
            if 2 * min(len(canon), len(rep_canon)) < LABEL_SIMILARITY * (len(canon) + len(rep_canon)):
                continue
            sm = SequenceMatcher(None, canon, rep_canon)
            if sm.quick_ratio() >= LABEL_SIMILARITY and sm.ratio() >= LABEL_SIMILARITY:
                target = rep
                break
        if target is None:
            reps.append((key, canon))
            target = key
        by_canon[canon] = target
        merged[key] = target
    return merged


def prepare(raw_rows, analyst, catalog, counts):
    """Aplica o tratamento e devolve (exemplos, contagens, classes próprias).

    Cada exemplo: {'text', 'raw', 'cls' (ou None), 'detail', 'weight', 'source'}."""
    counts = dict(counts)
    kept = []
    no_content = no_info = 0
    for raw, label, _ in raw_rows:
        norm = normalize_text(raw)
        if not has_content(norm):
            no_content += 1
            continue
        key = label_key(label)
        if not key or NO_INFO_LABEL.match(key):
            no_info += 1
            continue
        kept.append((norm, key, raw, label))
    counts.update(sem_conteudo=no_content, rotulo_sem_informacao=no_info)

    # Rótulos quase iguais viram um só; a grafia exibida é a mais usada no grupo.
    merged = merge_labels([k for _, k, _, _ in kept])
    counts['rotulos_antes'] = len({k for _, k, _, _ in kept})
    counts['rotulos_depois'] = len(set(merged.values()))
    spellings = defaultdict(Counter)
    for _, k, _, label in kept:
        spellings[merged[k]][label] += 1
    display = {g: c.most_common(1)[0][0] for g, c in spellings.items()}

    # Duplicadas e conflitos: um exemplo por relato.
    pairs = Counter((norm, merged[k]) for norm, k, _, _ in kept)
    counts['duplicadas'] = sum(n - 1 for n in pairs.values())
    first_raw = {}
    for norm, _, raw, _ in kept:
        first_raw.setdefault(norm, raw)
    by_text = defaultdict(Counter)
    for (norm, group), n in pairs.items():
        by_text[norm][group] = n
    global_freq = Counter()
    for (_, group), n in pairs.items():
        global_freq[group] += n
    conflicts = 0
    chosen = {}
    for norm, options in by_text.items():
        if len(options) > 1:
            conflicts += 1
        chosen[norm] = max(options, key=lambda g: (options[g], global_freq[g], g))
    counts['conflitos'] = conflicts

    # Exemplos do analista substituem o que a pasta dizia sobre o mesmo relato.
    analyst_by_text = {ex['relato_norm']: ex for ex in analyst if ex.get('relato_norm')}
    counts['substituidos_pelo_analista'] = sum(1 for norm in chosen if norm in analyst_by_text)

    group_size = Counter()
    for norm, group in chosen.items():
        if norm not in analyst_by_text:
            group_size[group] += sum(by_text[norm].values())
    unclassified = set(fc.UNCLASSIFIED)
    own_classes = set()
    examples = []
    mapped = Counter()
    for norm, group in chosen.items():
        if norm in analyst_by_text:
            continue
        label = display[group]
        cls = fc.classify(label, '', catalog)
        how = 'rotulo'
        if cls in unclassified:
            cls, how = fc.classify(first_raw[norm], '', catalog), 'relato'
        if cls in unclassified:
            if group_size[group] >= MIN_OWN_CLASS:
                cls, how = class_label(label), 'classe_propria'
                own_classes.add(cls)
            else:
                cls, how = None, 'so_detalhe'
        mapped[how] += 1
        examples.append({'text': norm, 'raw': first_raw[norm], 'cls': cls, 'detail': label,
                         'weight': sum(by_text[norm].values()), 'source': 'pasta'})
    for norm, ex in analyst_by_text.items():
        examples.append({'text': norm, 'raw': ex.get('relato') or norm, 'cls': class_label(ex['classe']),
                         'detail': ex.get('detalhe') or '', 'weight': ANALYST_WEIGHT, 'source': 'analista'})
    counts.update(classe_pelo_rotulo=mapped['rotulo'], classe_pelo_relato=mapped['relato'],
                  classe_propria=mapped['classe_propria'], so_detalhe=mapped['so_detalhe'],
                  do_analista=len(analyst_by_text), exemplos=len(examples),
                  exemplos_com_classe=sum(1 for e in examples if e['cls']))
    return examples, counts, sorted(own_classes)


# ----------------------------------------------------------------------------- modelo

def _build(C=1.0):
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.pipeline import make_pipeline, make_union
    from sklearn.svm import LinearSVC
    features = make_union(
        TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True),
        TfidfVectorizer(analyzer='char_wb', ngram_range=(3, 5), sublinear_tf=True),
    )
    return make_pipeline(features, LinearSVC(C=C))


def _fit(examples, C=1.0):
    model = _build(C)
    model.fit([e['text'] for e in examples], [e['cls'] for e in examples],
              linearsvc__sample_weight=[e['weight'] for e in examples])
    return model


def _proba(pipeline, texts):
    import numpy as np
    scores = pipeline.decision_function(texts)
    if scores.ndim == 1:  # só duas classes
        scores = np.column_stack([-scores, scores])
    e = np.exp((scores - scores.max(axis=1, keepdims=True)) / TEMPERATURE)
    return e / e.sum(axis=1, keepdims=True)


def _split(items, seed):
    index = list(range(len(items)))
    random.Random(seed).shuffle(index)
    cut = int(round(len(index) * (1 - TEST_SIZE)))
    return [items[i] for i in index[:cut]], [items[i] for i in index[cut:]]


def _accuracy(model, test):
    if not test:
        return 0.0
    pred = model.predict([e['text'] for e in test])
    return sum(p == e['cls'] for p, e in zip(pred, test)) / len(test)


class DetailIndex:
    """Vizinho mais próximo dentro da classe prevista → rótulo do analista desse vizinho."""

    def __init__(self, model, examples):
        self.features = model[:-1]
        pool = [e for e in examples if e['detail']]
        self.matrix = self.features.transform([e['text'] for e in pool]) if pool else None
        self.details = [e['detail'] for e in pool]
        self.by_class = defaultdict(list)
        for i, e in enumerate(pool):
            if e['cls']:
                self.by_class[e['cls']].append(i)

    def suggest(self, texts, classes):
        import numpy as np
        out = [''] * len(texts)
        if self.matrix is None or not texts:
            return out
        for start in range(0, len(texts), 1000):
            q = self.features.transform(texts[start:start + 1000])
            sims = (q @ self.matrix.T).toarray()
            for i, row in enumerate(sims):
                idx = self.by_class.get(classes[start + i])
                if idx:
                    out[start + i] = self.details[idx[int(np.argmax(row[idx]))]]
        return out

    def options(self, limit=40):
        result = {}
        for cls, idx in self.by_class.items():
            result[cls] = [d for d, _ in Counter(self.details[i] for i in idx).most_common(limit)]
        return result


def _wilson(hits, n, z=1.96):
    """Intervalo de confiança de 95% (Wilson) para uma proporção."""
    if not n:
        return 0.0, 0.0
    p = hits / n
    center = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return max(0.0, center - half), min(1.0, center + half)


def _pct(x):
    return round(x * 100, 1)


def band(confidence):
    return 'Alta' if confidence >= HIGH else 'Média' if confidence >= MEDIUM else 'Baixa'


def train(force=False):
    files = _training_files()
    analyst = _analyst_examples()
    fingerprint = _fingerprint(files, analyst)
    with _lock:
        if not force and _state['model'] is not None and _state['fingerprint'] == fingerprint:
            return _state['info']
        started = time.time()
        from . import service
        catalog = service.catalog()
        raw_rows, per_file, counts = read_folder()
        examples, cleaning, own_classes = prepare(raw_rows, analyst, catalog, counts)
        labeled = [e for e in examples if e['cls']]
        if len(labeled) < 50 or len({e['cls'] for e in labeled}) < 2:
            raise ValueError('Exemplos insuficientes para treinar (mínimo 50 relatos com classe).')

        folder = [e for e in labeled if e['source'] == 'pasta']
        reviewed = [e for e in labeled if e['source'] == 'analista']
        folder_train, test_set = _split(folder, SEED)
        train_set = folder_train + reviewed
        # Ajuste fino só dentro dos 80%: o teste continua intocado.
        inner_train, inner_val = _split(train_set, SEED + 1)
        scores = {C: _accuracy(_fit(inner_train, C), inner_val) for C in C_GRID}
        best_c = max(C_GRID, key=lambda C: (scores[C], -abs(C - 1)))

        model = _fit(train_set, best_c)
        in_train = {id(e) for e in train_set}
        detail_pool = [e for e in examples if id(e) in in_train or not e['cls']]
        details = DetailIndex(model, detail_pool)

        test_texts = [e['text'] for e in test_set]
        proba = _proba(model, test_texts)
        classes = model.classes_
        guesses = [classes[int(p.argmax())] for p in proba]
        suggested = details.suggest(test_texts, guesses)
        hits = top3 = detail_hits = 0
        bands = {b: [0, 0] for b in ('Alta', 'Média', 'Baixa')}
        errors = []
        for probs, guess, detail, e in zip(proba, guesses, suggested, test_set):
            order = probs.argsort()[::-1]
            conf = float(probs[order[0]])
            ok = guess == e['cls']
            hits += ok
            top3 += e['cls'] in {classes[j] for j in order[:3]}
            detail_hits += bool(e['detail']) and label_key(detail) == label_key(e['detail'])
            b = bands[band(conf)]
            b[0] += 1
            b[1] += ok
            if not ok and len(errors) < 12:
                errors.append({'relato': e['raw'], 'esperado': e['cls'], 'previsto': guess, 'confianca': _pct(conf)})
        n = len(test_set)
        low, high = _wilson(hits, n)
        accuracy = hits / n

        # Estabilidade: o mesmo processo com outros sorteios 80/20.
        runs = [accuracy] + [_accuracy(_fit(tr + reviewed, best_c), te) for tr, te in
                             (_split(folder, s) for s in STABILITY_SEEDS if s != SEED)]
        mean = sum(runs) / len(runs)
        sd = math.sqrt(sum((r - mean) ** 2 for r in runs) / len(runs))

        previous = (_state['info'] or {}).get('accuracy')
        info = {
            'ready': True,
            'trained_at': time.strftime('%Y-%m-%dT%H:%M:%S'),
            'seconds': round(time.time() - started, 1),
            'train_dir': str(train_dir()),
            'files': per_file,
            'cleaning': cleaning,
            'examples': len(labeled),
            'train_size': len(train_set),
            'test_size': n,
            'split': '80/20',
            'C': best_c,
            'classes': len(classes),
            'own_classes': own_classes,
            'accuracy': _pct(accuracy),
            'previous_accuracy': previous,
            'error_rate': _pct(1 - accuracy),
            'margin': _pct((high - low) / 2),
            'ci_low': _pct(low),
            'ci_high': _pct(high),
            'stability_mean': _pct(mean),
            'stability_sd': _pct(sd),
            'stability_runs': len(runs),
            'top3': _pct(top3 / n),
            'detail_accuracy': _pct(detail_hits / n),
            'analyst_examples': cleaning['do_analista'],
            'bands': [
                {'band': b, 'min': {'Alta': HIGH, 'Média': MEDIUM, 'Baixa': 0}[b] * 100,
                 'share': _pct(v[0] / n), 'accuracy': _pct(v[1] / v[0]) if v[0] else None, 'count': v[0]}
                for b, v in bands.items()
            ],
            'error_examples': errors,
            'top_classes': [{'label': k, 'examples': c} for k, c in Counter(e['cls'] for e in labeled).most_common(10)],
        }
        bundle = {'pipeline': model, 'details': details,
                  'memory': {ex['relato_norm']: ex for ex in analyst if ex.get('relato_norm')},
                  'class_list': sorted(set(classes) | {e['cls'] for e in reviewed} | _catalog_classes(catalog))}
        _state.update(model=bundle, info=info, fingerprint=fingerprint)
        try:
            import joblib
            joblib.dump({'model': bundle, 'info': info, 'fingerprint': fingerprint}, _model_path())
        except Exception:
            pass
        return info


def _catalog_classes(catalog):
    # Componentes viram "FALHA DE <componente>", como no quadro.
    return {f"FALHA DE {c['label'].strip().upper()}" if c.get('kind', 'componente') == 'componente'
            else c['label'].strip().upper() for c in catalog.get('classes', []) if c.get('label')}


def _current_fingerprint():
    try:
        return _fingerprint(_training_files(), _analyst_examples())
    except ValueError:
        return None


def _ensure_model():
    if _state['model'] is None:
        try:
            import joblib
            saved = joblib.load(_model_path())
            if 'details' in saved['model']:
                _state.update(model=saved['model'], info=saved['info'], fingerprint=saved['fingerprint'])
        except Exception:
            pass
    current = _current_fingerprint()
    # Retreina sozinho se a pasta de treino ou as revisões do analista mudaram.
    if _state['model'] is None or (current and current != _state['fingerprint']):
        train(force=True)
    return _state['model']


def retrain_soon(delay=20.0):
    """Retreina ~20 s depois da última revisão salva (cada nova revisão adia), sem travar a tela."""
    def run():
        if _state['training']:
            return
        _state['training'] = True
        try:
            train()
        except Exception:
            pass
        finally:
            _state['training'] = False
    if _timer['t']:
        _timer['t'].cancel()
    _timer['t'] = threading.Timer(delay, run)
    _timer['t'].daemon = True
    _timer['t'].start()


def status():
    try:
        if not _state['training']:
            _ensure_model()
        info = dict(_state['info'] or {})
        info['training'] = _state['training']
        info['pending'] = bool(_state['fingerprint'] and _current_fingerprint() != _state['fingerprint'])
        return info
    except ValueError as exc:
        return {'ready': False, 'erro': str(exc), 'train_dir': str(train_dir())}


def class_options():
    model = _ensure_model()
    return {'classes': model['class_list'], 'details': model['details'].options()}


# ----------------------------------------------------------------------------- predição

def predict(texts):
    """relato normalizado → {'classe', 'detalhe', 'confianca', 'faixa', 'top3'}."""
    model = _ensure_model()
    pipeline, memory = model['pipeline'], model['memory']
    out = {}
    unique = sorted({t for t in texts if t})
    for t in unique:
        if t in memory:  # relato que o analista já revisou: vale o que ele disse
            ex = memory[t]
            out[t] = {'classe': class_label(ex['classe']), 'detalhe': ex.get('detalhe') or '', 'confianca': None,
                      'faixa': ANALYST, 'top3': []}
    rest = [t for t in unique if t not in out]
    if rest:
        proba = _proba(pipeline, rest)
        classes = pipeline.classes_
        guesses = [classes[int(p.argmax())] for p in proba]
        details = model['details'].suggest(rest, guesses)
        for t, probs, guess, detail in zip(rest, proba, guesses, details):
            order = probs.argsort()[::-1][:3]
            conf = float(probs[order[0]])
            out[t] = {'classe': guess, 'detalhe': detail, 'confianca': _pct(conf), 'faixa': band(conf),
                      'top3': [{'classe': classes[j], 'confianca': _pct(float(probs[j]))} for j in order]}
    return out


FILLS = {'Alta': 'E8F6EF', 'Média': 'FDF3E2', 'Baixa': 'FDECEE', ANALYST: 'EBF1FC'}


def _paint(ws, row, cols, item):
    from openpyxl.styles import PatternFill
    fill = PatternFill('solid', fgColor=FILLS.get(item['faixa'], 'FFFFFF'))
    # .value explícito: ws.cell(r, c, None) não apaga o valor anterior.
    for key, value in (('label', item['classe']), ('detail', item['detalhe'] or None),
                       ('conf', item['confianca']), ('band', item['faixa'])):
        ws.cell(row, cols[key]).value = value
    ws.cell(row, cols['label']).fill = fill
    ws.cell(row, cols['band']).fill = fill


def fill_workbook(data, filename):
    """Preenche a planilha enviada e devolve (id do arquivo gerado, resumo para a tela)."""
    import openpyxl
    from openpyxl.styles import Font

    ext = Path(filename).suffix.lower()
    if ext not in EXTENSIONS:
        raise ValueError('Envie uma planilha .xlsx ou .xlsm.')
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=ext == '.xlsm')
    except Exception as exc:
        raise ValueError('Não foi possível ler o arquivo. Confira se é uma planilha válida, sem senha.') from exc

    target = None
    for ws in wb.worksheets:
        head = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=1, max_row=30)]
        header, text_col, label_col = _find_header(head)
        if header is not None:
            target = (ws, header + 1, text_col + 1, label_col + 1 if label_col is not None else None)
            break
    if not target:
        raise ValueError('Nenhuma aba tem a coluna "Observações" nas primeiras 30 linhas.')
    ws, header_row, text_col, label_col = target

    last = ws.max_column
    bold = Font(bold=True)
    if label_col is None:
        last += 1
        label_col = last
        ws.cell(header_row, label_col, TARGET).font = bold
    cols = {'label': label_col, 'detail': last + 1, 'conf': last + 2, 'band': last + 3}
    for key, title in (('detail', DETAIL_COLUMN), ('conf', CONF_COLUMN), ('band', BAND_COLUMN)):
        ws.cell(header_row, cols[key], title).font = bold

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(r, text_col).value
        if raw is None or str(raw).strip() == '':
            continue
        rows.append((r, str(raw), normalize_text(raw), ws.cell(r, label_col).value))
    if not rows:
        raise ValueError('A coluna "Observações" está vazia.')

    predictions = predict([n for _, _, n, existing in rows if existing in (None, '') and has_content(n)])
    counts = Counter()
    labels = Counter()
    items = {}
    preview = []
    for r, raw, norm, existing in rows:
        if existing not in (None, ''):
            counts['Preenchida'] += 1
            ws.cell(r, cols['band'], 'Preenchida antes')
            continue
        if not has_content(norm):
            counts['Sem relato'] += 1
            ws.cell(r, cols['band'], 'Sem relato para classificar')
            continue
        p = predictions[norm]
        counts[p['faixa']] += 1
        labels[p['classe']] += 1
        _paint(ws, r, cols, p)
        item = items.get(norm)
        if item is None:
            item = items[norm] = {'key': norm, 'relato': raw, 'linhas': [], **p,
                                  'status': 'revisado' if p['faixa'] == ANALYST else 'pendente'}
        item['linhas'].append(r)
        if len(preview) < 300:
            preview.append({'linha': r, 'relato': raw, **{k: p[k] for k in ('classe', 'detalhe', 'confianca', 'faixa')},
                            'alternativas': p['top3'][1:]})
    for key, width in (('label', 34), ('detail', 40), ('band', 18)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cols[key])].width = width

    out = io.BytesIO()
    wb.save(out)
    file_id = uuid4().hex
    out_ext = '.xlsm' if ext == '.xlsm' else '.xlsx'
    (_outputs_dir() / f'{file_id}{out_ext}').write_bytes(out.getvalue())
    info = _state['info']
    run = {
        'id': file_id,
        'filename': f'{Path(filename).stem} - classificada{out_ext}',
        'source': filename,
        'created_at': time.strftime('%Y-%m-%dT%H:%M:%S'),
        'sheet': ws.title,
        'cols': cols,
        'rows': len(rows),
        'dirty': False,
        'items': list(items.values()),
    }
    _write_run(run)
    summary = {
        'id': file_id,
        'filename': run['filename'],
        'sheet': ws.title,
        'rows': len(rows),
        'filled': sum(counts[b] for b in ('Alta', 'Média', 'Baixa', ANALYST)),
        'unique': len(items),
        'counts': {k: counts.get(k, 0) for k in ('Alta', 'Média', 'Baixa', ANALYST, 'Preenchida', 'Sem relato')},
        # Acerto esperado: acurácia de cada faixa no teste × linhas da planilha em cada faixa
        # (revisadas pelo analista contam como certas).
        'expected_accuracy': _expected(counts, info),
        'top_labels': [{'label': k, 'rows': v} for k, v in labels.most_common(10)],
        'preview': preview,
        'model': info,
    }
    return file_id, summary


def _expected(counts, info):
    by_band = {b['band']: b['accuracy'] or 0 for b in info['bands']}
    by_band[ANALYST] = 100
    total = sum(counts[b] for b in by_band)
    if not total:
        return None
    return round(sum(counts[b] * by_band[b] for b in by_band) / total, 1)


# ----------------------------------------------------------------------------- revisão do analista

def _run_path(file_id):
    if not re.fullmatch(r'[0-9a-f]{32}', file_id or ''):
        return None
    return _outputs_dir() / f'{file_id}.json'


def _write_run(run):
    _run_path(run['id']).write_text(json.dumps(run, ensure_ascii=False), encoding='utf-8')


def read_run(file_id):
    path = _run_path(file_id)
    if not path or not path.exists():
        raise ValueError('Planilha processada não encontrada. Gere a planilha novamente.')
    return json.loads(path.read_text(encoding='utf-8'))


def _run_summary(run):
    done = sum(1 for i in run['items'] if i['status'] != 'pendente')
    return {'id': run['id'], 'filename': run['filename'], 'created_at': run['created_at'], 'rows': run['rows'],
            'unique': len(run['items']), 'reviewed': done,
            'pending_low': sum(1 for i in run['items'] if i['status'] == 'pendente' and i['faixa'] != 'Alta')}


def list_runs():
    runs = []
    for path in sorted(_outputs_dir().glob('*.json'), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            runs.append(_run_summary(json.loads(path.read_text(encoding='utf-8'))))
        except Exception:
            continue
    return runs


ORDER = {'Baixa': 0, 'Média': 1, 'Alta': 2, ANALYST: 3}


def run_items(file_id, faixa='', status='', q='', page=1, size=100):
    run = read_run(file_id)
    items = run['items']
    if faixa:
        items = [i for i in items if i['faixa'] == faixa]
    if status:
        items = [i for i in items if i['status'] == status]
    if q:
        term = _plain(q)
        items = [i for i in items if term in i['key'] or term in _plain(i['classe']) or term in _plain(i['detalhe'])]
    # O que mais precisa de olho primeiro: pendente, menor confiança, mais linhas na planilha.
    items = sorted(items, key=lambda i: (i['status'] != 'pendente', ORDER.get(i['faixa'], 3),
                                         i['confianca'] if i['confianca'] is not None else 100, -len(i['linhas'])))
    total = len(items)
    start = (max(page, 1) - 1) * size
    page_items = [{**{k: v for k, v in i.items() if k != 'linhas'}, 'n_linhas': len(i['linhas']),
                   'linhas': i['linhas'][:5]} for i in items[start:start + size]]
    return {'run': _run_summary(run), 'total': total, 'page': max(page, 1), 'size': size, 'items': page_items}


def review(file_id, entries):
    """Grava confirmações/correções do analista: viram exemplos de treino e corrigem a planilha gerada."""
    from . import store
    if not isinstance(entries, list) or not entries:
        raise ValueError('Nenhuma revisão para salvar.')
    with _runs_lock:
        run = read_run(file_id)
        by_key = {i['key']: i for i in run['items']}
        saved = []
        for entry in entries[:1000]:
            item = by_key.get(str((entry or {}).get('key') or ''))
            if not item:
                continue
            action = entry.get('acao')
            if action == 'confirmar':
                classe, detalhe = item['classe'], item['detalhe']
            elif action == 'corrigir':
                classe = class_label(entry.get('classe'))
                detalhe = re.sub(r'\s+', ' ', str(entry.get('detalhe') or '')).strip()[:300]
                if not classe:
                    raise ValueError('Informe a classe correta.')
            else:
                raise ValueError('Ação de revisão desconhecida.')
            item.update(classe=classe, detalhe=detalhe, faixa=ANALYST, confianca=None,
                        status='confirmado' if action == 'confirmar' else 'corrigido')
            saved.append({'relato': item['relato'], 'relato_norm': item['key'], 'classe': classe,
                          'detalhe': detalhe, 'origem': run['source']})
        if not saved:
            raise ValueError('Os relatos enviados não pertencem a esta planilha.')
        store.save_ml_examples(saved)
        run['dirty'] = True  # a planilha é reescrita na hora do download
        _write_run(run)
    retrain_soon()
    return {'saved': len(saved), 'run': _run_summary(run)}


def output_path(file_id):
    """Arquivo pronto para download, já com as revisões do analista aplicadas."""
    if not re.fullmatch(r'[0-9a-f]{32}', file_id or ''):
        return None
    path = next((p for p in (_outputs_dir() / f'{file_id}{e}' for e in EXTENSIONS) if p.exists()), None)
    if not path:
        return None
    with _runs_lock:
        try:
            run = read_run(file_id)
        except ValueError:
            return path
        if run.get('dirty'):
            import openpyxl
            wb = openpyxl.load_workbook(path, keep_vba=path.suffix == '.xlsm')
            ws = wb[run['sheet']]
            for item in run['items']:
                if item['status'] in ('confirmado', 'corrigido'):
                    for r in item['linhas']:
                        _paint(ws, r, run['cols'], item)
            wb.save(path)
            run['dirty'] = False
            _write_run(run)
    return path
